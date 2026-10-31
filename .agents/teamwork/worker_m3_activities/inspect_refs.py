import os
from docx import Document

def inspect_docx(path):
    print("=" * 60)
    print("FILE:", path)
    if not os.path.exists(path):
        print("DOES NOT EXIST")
        return
    try:
        doc = Document(path)
    except Exception as e:
        print("Error opening doc:", e)
        return
    print("Paragraphs count:", len(doc.paragraphs))
    print("Tables count:", len(doc.tables))
    print("Sections count:", len(doc.sections))
    for i, s in enumerate(doc.sections):
        try:
            print(f"Section {i}: top={s.top_margin.inches:.2f}\", bottom={s.bottom_margin.inches:.2f}\", left={s.left_margin.inches:.2f}\", right={s.right_margin.inches:.2f}\", orientation={s.orientation}, page_w={s.page_width.inches:.2f}\", page_h={s.page_height.inches:.2f}\"")
        except Exception as e:
            print(f"Section {i} margin read error: {e}")
    
    print("\nPARAGRAPH HEADINGS & SAMPLES:")
    for i, p in enumerate(doc.paragraphs):
        txt = p.text.strip()
        if txt and (len(txt) < 80 or p.style.name.startswith('Heading') or 'Part' in txt or 'Activity' in txt or 'Table' in txt or 'Figure' in txt):
            print(f"P{i:03d} [{p.style.name}]: {txt}")
            
    print("\nTABLE SUMMARY:")
    for i, t in enumerate(doc.tables):
        rows = len(t.rows)
        cols = len(t.columns) if rows > 0 else 0
        first_row = [c.text.strip().replace('\n', ' ') for c in t.rows[0].cells] if rows > 0 else []
        print(f"Table {i+1}: {rows} rows, {cols} cols | Headers: {first_row[:6]}")

if __name__ == '__main__':
    inspect_docx(r"c:\Users\Admin\Downloads\Activity#2_DFD (3).docx")
    inspect_docx(r"c:\Users\Admin\Downloads\Activity#4_UML_Object.docx")
    inspect_docx(r"c:\Users\Admin\AndroidStudioProjects\class_realm01\Capstone-1-and-Thesis-1-Document-FORMATS-1785742696-1.docx")
