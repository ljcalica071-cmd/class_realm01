import docx

doc = docx.Document(r"c:\Users\Admin\AndroidStudioProjects\class_realm01\INTEGRATED_SENIOR_HIGH_SCHOOL_ATTENDANCE_CH1_CH2_CH3.docx")

print("Checking paragraphs for drawings:")
for idx, p in enumerate(doc.paragraphs):
    if "drawing" in p._p.xml:
        print(f"Para {idx}: text='{p.text}', xml snippet={p._p.xml[:120]}")
        # check next 2 paragraphs
        for next_i in range(idx+1, min(idx+3, len(doc.paragraphs))):
            print(f"   Next Para {next_i}: '{doc.paragraphs[next_i].text}'")
