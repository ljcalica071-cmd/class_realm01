import os
import sys
import docx
from docx.oxml import OxmlElement
from docx.oxml.ns import qn, nsdecls
import xml.etree.ElementTree as ET

DOCX_PATH = r"c:\Users\Admin\AndroidStudioProjects\class_realm01\INTEGRATED_SENIOR_HIGH_SCHOOL_ATTENDANCE_CH1_CH2_CH3.docx"

print(f"=== AUDITING: {DOCX_PATH} ===")
if not os.path.exists(DOCX_PATH):
    print("ERROR: File does not exist!")
    sys.exit(1)

stat = os.stat(DOCX_PATH)
print(f"File size: {stat.st_size:,} bytes")
print(f"Last modified: {stat.st_mtime}")

doc = docx.Document(DOCX_PATH)

print("\n--- SECTIONS & MARGINS ---")
for idx, sec in enumerate(doc.sections):
    print(f"Section {idx}:")
    print(f"  Top: {sec.top_margin.inches:.2f} in ({sec.top_margin.pt:.1f} pt)")
    print(f"  Bottom: {sec.bottom_margin.inches:.2f} in ({sec.bottom_margin.pt:.1f} pt)")
    print(f"  Left: {sec.left_margin.inches:.2f} in ({sec.left_margin.pt:.1f} pt)")
    print(f"  Right: {sec.right_margin.inches:.2f} in ({sec.right_margin.pt:.1f} pt)")
    print(f"  Page Size: {sec.page_width.inches:.2f} x {sec.page_height.inches:.2f} in")

print("\n--- TABLES AUDIT ---")
print(f"Total tables: {len(doc.tables)}")
for t_idx, table in enumerate(doc.tables):
    rows_count = len(table.rows)
    cols_count = len(table.columns)
    cant_split_count = 0
    tbl_header_count = 0
    for r_idx, row in enumerate(table.rows):
        trPr = row._tr.find(qn('w:trPr'))
        if trPr is not None:
            if trPr.find(qn('w:cantSplit')) is not None:
                cant_split_count += 1
            if trPr.find(qn('w:tblHeader')) is not None:
                tbl_header_count += 1
    
    # preview first row
    first_row_txt = [c.text.strip().replace('\n', ' ')[:30] for c in table.rows[0].cells]
    print(f"Table {t_idx}: {rows_count} rows, {cols_count} cols | cantSplit: {cant_split_count}/{rows_count} rows | tblHeader: {tbl_header_count} | Header preview: {first_row_txt}")

print("\n--- PARAGRAPHS & HEADINGS SUMMARY ---")
headings = []
for p_idx, p in enumerate(doc.paragraphs):
    txt = p.text.strip()
    if not txt:
        continue
    if p.style.name.startswith('Heading') or txt.isupper() or any(k in txt for k in ['CHAPTER', 'Project Context', 'Purpose and Description', 'Objectives', 'Scope and Limitation', 'Definition of Terms', 'Theoretical Framework', 'Conceptual Framework', 'Methodology', 'Table ', 'Figure ']):
        if len(txt) < 80:
            headings.append((p_idx, p.style.name, txt))

print(f"Total paragraphs: {len(doc.paragraphs)}")
print(f"Sample headings/key anchors ({len(headings)} found):")
for p_idx, style, txt in headings[:35]:
    print(f"  [{p_idx}] ({style}): {txt}")

print("\n--- PAGE BREAKS & CHAPTER 1 BOUNDARY AUDIT ---")
# Count page breaks in body
page_breaks = []
ch1_start_p = None
ch1_end_p = None
ch2_start_p = None
ch3_start_p = None

for p_idx, p in enumerate(doc.paragraphs):
    txt = p.text.strip()
    # Check for w:br type="page"
    for r in p._p.findall(qn('w:r')):
        for br in r.findall(qn('w:br')):
            if br.get(qn('w:type')) == 'page':
                page_breaks.append((p_idx, 'page_break_run', txt[:30]))
    # Check for paragraph pageBreakBefore
    pPr = p._p.find(qn('w:pPr'))
    if pPr is not None and pPr.find(qn('w:pageBreakBefore')) is not None:
        page_breaks.append((p_idx, 'pageBreakBefore', txt[:30]))

    if 'CHAPTER I' in txt.upper() and ch1_start_p is None:
        ch1_start_p = p_idx
    elif 'CHAPTER II' in txt.upper() and ch2_start_p is None:
        ch2_start_p = p_idx
    elif 'CHAPTER III' in txt.upper() and ch3_start_p is None:
        ch3_start_p = p_idx

print(f"Total explicit page breaks found in XML: {len(page_breaks)}")
for pb in page_breaks:
    print(f"  At p_idx {pb[0]}: {pb[1]} near '{pb[2]}'")

print(f"\nAnchors: Chapter 1 start p_idx={ch1_start_p}, Chapter 2 start p_idx={ch2_start_p}, Chapter 3 start p_idx={ch3_start_p}")
