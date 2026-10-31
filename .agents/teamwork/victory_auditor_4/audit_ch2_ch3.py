import docx
import re
import sys

sys.stdout.reconfigure(encoding='utf-8', errors='replace')

DOCX_PATH = r"c:\Users\Admin\AndroidStudioProjects\class_realm01\INTEGRATED_SENIOR_HIGH_SCHOOL_ATTENDANCE_CH1_CH2_CH3.docx"
doc = docx.Document(DOCX_PATH)

print("=== AUDIT CHAPTER 2 & CHAPTER 3 ===")

# 1. Inspect Table 1 (Literature review matrix)
t1 = doc.tables[0]
print(f"\n--- Table 1 Review ---")
print(f"Table 1 dimensions: {len(t1.rows)} rows x {len(t1.columns)} cols")
header = [c.text.strip().replace('\n', ' ') for c in t1.rows[0].cells]
print("Header:", header)

studies = []
for r_idx in range(1, len(t1.rows)):
    row = t1.rows[r_idx]
    author = row.cells[0].text.strip().replace('\n', ' ')
    title_venue = row.cells[1].text.strip().replace('\n', ' ')
    method = row.cells[2].text.strip().replace('\n', ' ')
    sample = row.cells[3].text.strip().replace('\n', ' ')
    findings = row.cells[4].text.strip().replace('\n', ' ')
    contrib = row.cells[5].text.strip().replace('\n', ' ')
    studies.append({
        'idx': r_idx,
        'author': author,
        'title': title_venue,
        'method': method,
        'sample': sample,
        'findings': findings,
        'contrib': contrib
    })
    print(f"\nStudy {r_idx}: {author}")
    print(f"  Title/Venue: {title_venue[:80]}...")
    print(f"  Method: {method[:60]}...")
    print(f"  Sample: {sample[:60]}...")
    print(f"  Findings: {findings[:80]}...")

print(f"\nTotal studies in Table 1: {len(studies)}")

# Check years (2021-2026)
years = []
for s in studies:
    match = re.search(r'202[1-6]', s['author'] + ' ' + s['title'])
    years.append(match.group(0) if match else "UNKNOWN")
print("Study years found:", years)

# Check for blanket "lacks" statements
lacks_count = 0
for s in studies:
    text_all = (s['findings'] + ' ' + s['contrib']).lower()
    if 'lacks' in text_all or 'lack of' in text_all:
        print(f"Warning: 'lack' found in study {s['idx']}: {s['author']}")
        lacks_count += 1
print(f"Total 'lacks' occurrences in Table 1: {lacks_count}")

# 2. Check DeLone & McLean
print("\n--- Theoretical Framework & DeLone & McLean ---")
dm_found = []
for p_idx, p in enumerate(doc.paragraphs[100:172]):
    txt = p.text.strip()
    if 'delone' in txt.lower() or 'mclean' in txt.lower() or 'information systems success' in txt.lower():
        dm_found.append((p_idx + 100, txt))
for idx, txt in dm_found:
    print(f"P{idx}: {txt[:120]}...")

# 3. Check Conceptual Framework (IPSO & IPO)
print("\n--- Conceptual Framework & Figures ---")
ipso_ipo_found = []
for p_idx, p in enumerate(doc.paragraphs[100:172]):
    txt = p.text.strip()
    if any(k in txt for k in ['IPSO', 'IPO', 'Figure 1', 'Figure 2', 'Input-Process']):
        ipso_ipo_found.append((p_idx + 100, txt))
for idx, txt in ipso_ipo_found:
    print(f"P{idx}: {txt[:120]}...")

# 4. Check Chapter 3 Methodology & Purposive Sampling
print("\n--- Chapter 3 Methodology & Purposive Sampling ---")
sampling_found = []
for p_idx, p in enumerate(doc.paragraphs[172:]):
    txt = p.text.strip()
    if any(k in txt.lower() for k in ['purposive', 'karen', 'gonzaga', 'teacher', 'parent', 'scaling', 'multi-teacher']):
        sampling_found.append((p_idx + 172, txt))
print(f"Found {len(sampling_found)} paragraphs referencing sampling/teacher/parent/scaling in Ch3:")
for idx, txt in sampling_found[:10]:
    print(f"P{idx}: {txt[:120]}...")

# 5. Check Figures in Chapter 3 (Figures 3, 4, 5)
print("\n--- Chapter 3 Figures ---")
for p_idx, p in enumerate(doc.paragraphs[172:]):
    txt = p.text.strip()
    if 'figure ' in txt.lower():
        print(f"P{p_idx + 172}: {txt}")
