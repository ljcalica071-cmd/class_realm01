import docx

doc = docx.Document(r"c:\Users\Admin\AndroidStudioProjects\class_realm01\Capstone-1-and-Thesis-1-Document-FORMATS-1785742696-1.docx")

print("=== PARAGRAPHS IN FORMAT GUIDELINES ===")
for i, p in enumerate(doc.paragraphs):
    if p.text.strip():
        print(f"P{i}: {p.text}")

print("\n=== TABLES IN FORMAT GUIDELINES ===")
for t_idx, table in enumerate(doc.tables):
    print(f"\n--- Table {t_idx} ---")
    for row in table.rows:
        row_text = [cell.text.strip().replace('\n', ' ') for cell in row.cells]
        print(" | ".join(row_text))

print("\n=== SECTIONS INFO ===")
for s_idx, section in enumerate(doc.sections):
    print(f"Section {s_idx}:")
    print(f"  Top margin: {section.top_margin.inches} in ({section.top_margin.pt} pt)")
    print(f"  Bottom margin: {section.bottom_margin.inches} in ({section.bottom_margin.pt} pt)")
    print(f"  Left margin: {section.left_margin.inches} in ({section.left_margin.pt} pt)")
    print(f"  Right margin: {section.right_margin.inches} in ({section.right_margin.pt} pt)")
    print(f"  Page width: {section.page_width.inches} in, height: {section.page_height.inches} in")
