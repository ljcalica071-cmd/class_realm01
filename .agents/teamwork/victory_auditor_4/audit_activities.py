import os
import sys
import docx
from docx.oxml.ns import qn

sys.stdout.reconfigure(encoding='utf-8', errors='replace')

DOCX_PATH = r"c:\Users\Admin\AndroidStudioProjects\class_realm01\CLASSREALM_Activities_2_and_4.docx"

print(f"=== AUDITING DELIVERABLE 2: {DOCX_PATH} ===")
if not os.path.exists(DOCX_PATH):
    print("ERROR: File does not exist!")
    sys.exit(1)

stat = os.stat(DOCX_PATH)
print(f"File size: {stat.st_size:,} bytes")
print(f"Last modified: {stat.st_mtime}")

doc = docx.Document(DOCX_PATH)

print("\n--- SECTIONS & MARGINS ---")
for idx, sec in enumerate(doc.sections):
    print(f"Section {idx}:")
    print(f"  Top: {sec.top_margin.inches:.2f} in ({sec.top_margin.pt:.1f} pt)")
    print(f"  Bottom: {sec.bottom_margin.inches:.2f} in ({sec.bottom_margin.pt:.1f} pt)")
    print(f"  Left: {sec.left_margin.inches:.2f} in ({sec.left_margin.pt:.1f} pt)")
    print(f"  Right: {sec.right_margin.inches:.2f} in ({sec.right_margin.pt:.1f} pt)")
    print(f"  Page Size: {sec.page_width.inches:.2f} x {sec.page_height.inches:.2f} in")

print("\n--- TABLES AUDIT ---")
print(f"Total tables: {len(doc.tables)}")
for t_idx, table in enumerate(doc.tables):
    rows_count = len(table.rows)
    cols_count = len(table.columns)
    cant_split_count = 0
    tbl_header_count = 0
    for r_idx, row in enumerate(table.rows):
        trPr = row._tr.find(qn('w:trPr'))
        if trPr is not None:
            if trPr.find(qn('w:cantSplit')) is not None:
                cant_split_count += 1
            if trPr.find(qn('w:tblHeader')) is not None:
                tbl_header_count += 1
    
    first_row_txt = [c.text.strip().replace('\n', ' ')[:25] for c in table.rows[0].cells]
    print(f"Table {t_idx}: {rows_count} rows, {cols_count} cols | cantSplit: {cant_split_count}/{rows_count} | tblHeader: {tbl_header_count} | Header: {first_row_txt}")

print("\n--- HEADINGS & STRUCTURE ---")
act2_found = False
act4_found = False
headings = []
for p_idx, p in enumerate(doc.paragraphs):
    txt = p.text.strip()
    if not txt:
        continue
    if any(k in txt.upper() for k in ['ACTIVITY #2', 'ACTIVITY 2', 'ACTIVITY #4', 'ACTIVITY 4', 'DATA FLOW', 'CONTEXT DIAGRAM', 'LEVEL 1', 'ERD', 'USE CASE', 'CLASS DIAGRAM', 'OBJECT DIAGRAM', 'TABLE ', 'FIGURE ']):
        if len(txt) < 80:
            headings.append((p_idx, txt))
    if 'ACTIVITY #2' in txt.upper() or 'ACTIVITY 2' in txt.upper():
        act2_found = True
    if 'ACTIVITY #4' in txt.upper() or 'ACTIVITY 4' in txt.upper():
        act4_found = True

print(f"Activity #2 present: {act2_found}")
print(f"Activity #4 present: {act4_found}")
print(f"Key headings ({len(headings)}):")
for p_idx, txt in headings:
    print(f"  P{p_idx}: {txt}")

# Check Tech Stack mentions in document
print("\n--- TECH STACK AUDIT ---")
full_text = " ".join([p.text for p in doc.paragraphs])
for t in doc.tables:
    for r in t.rows:
        full_text += " " + " ".join([c.text for c in r.cells])

full_lower = full_text.lower()
tech_terms = ['drift', 'sqlite', 'esp32', 'rc522', 'supabase', 'firebase auth', 'mysql', 'mongodb', 'arduino uno', 'node.js']
for term in tech_terms:
    count = full_lower.count(term)
    print(f"Term '{term}': {count} occurrences")
