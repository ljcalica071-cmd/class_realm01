import docx
import sys
sys.stdout.reconfigure(encoding='utf-8', errors='replace')

DOCX_PATH = r"c:\Users\Admin\AndroidStudioProjects\class_realm01\INTEGRATED_SENIOR_HIGH_SCHOOL_ATTENDANCE_CH1_CH2_CH3.docx"
doc = docx.Document(DOCX_PATH)

for i in range(188, 198):
    if i < len(doc.paragraphs):
        print(f"--- P{i} ---")
        print(doc.paragraphs[i].text)
