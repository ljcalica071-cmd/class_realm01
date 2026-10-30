"""
CLASSREALM Chapters 1 and 2 Compiler Script
Injects complete, rigorous academic text and figures into INTEGRATED SENIOR HIGH SCHOOL ATTENDANCE.docx
Preserving Title Page front matter (elements 0 to 27) and replacing elements 28 to 50.
"""

import os
import sys
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

DOC_PATH = r"c:\Users\Admin\AndroidStudioProjects\class_realm01\INTEGRATED SENIOR HIGH SCHOOL ATTENDANCE.docx"
FIG1_PATH = r"c:\Users\Admin\AndroidStudioProjects\class_realm01\.agents\teamwork\worker_ch1_ch2_compiler\figures\image1.png"
FIG2_PATH = r"c:\Users\Admin\AndroidStudioProjects\class_realm01\.agents\teamwork\worker_ch1_ch2_compiler\figures\image2.png"

def apply_p_format(p, first_line_indent=0.5, line_spacing=1.5, space_before=0, space_after=0, align=WD_ALIGN_PARAGRAPH.LEFT, keep_with_next=False):
    if first_line_indent is not None:
        p.paragraph_format.first_line_indent = Inches(first_line_indent)
    else:
        p.paragraph_format.first_line_indent = None
    p.paragraph_format.line_spacing = line_spacing
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    p.alignment = align
    p.paragraph_format.keep_with_next = keep_with_next

def add_chapter_title(doc, chapter_num, chapter_title):
    p1 = doc.add_paragraph()
    apply_p_format(p1, first_line_indent=0, line_spacing=1.15, space_before=12, space_after=0, align=WD_ALIGN_PARAGRAPH.CENTER, keep_with_next=True)
    r1 = p1.add_run(chapter_num)
    r1.font.name = "Arial"
    r1.font.size = Pt(12)
    r1.font.bold = True

    p2 = doc.add_paragraph()
    apply_p_format(p2, first_line_indent=0, line_spacing=1.15, space_before=0, space_after=18, align=WD_ALIGN_PARAGRAPH.CENTER, keep_with_next=True)
    r2 = p2.add_run(chapter_title)
    r2.font.name = "Arial"
    r2.font.size = Pt(12)
    r2.font.bold = True

def add_h1(doc, text):
    p = doc.add_paragraph()
    apply_p_format(p, first_line_indent=0, line_spacing=1.5, space_before=16, space_after=2, align=WD_ALIGN_PARAGRAPH.LEFT, keep_with_next=True)
    r = p.add_run(text)
    r.font.name = "Arial"
    r.font.size = Pt(12)
    r.font.bold = True
    return p

def add_h2(doc, text):
    p = doc.add_paragraph()
    apply_p_format(p, first_line_indent=0, line_spacing=1.5, space_before=10, space_after=2, align=WD_ALIGN_PARAGRAPH.LEFT, keep_with_next=True)
    r = p.add_run(text)
    r.font.name = "Arial"
    r.font.size = Pt(11)
    r.font.bold = True
    return p

def add_h3(doc, text):
    p = doc.add_paragraph()
    apply_p_format(p, first_line_indent=0, line_spacing=1.5, space_before=6, space_after=2, align=WD_ALIGN_PARAGRAPH.LEFT, keep_with_next=True)
    r = p.add_run(text)
    r.font.name = "Arial"
    r.font.size = Pt(11)
    r.font.bold = True
    r.font.italic = True
    return p

def add_body_p(doc, text, bold_prefix=None):
    p = doc.add_paragraph()
    apply_p_format(p, first_line_indent=0.5, line_spacing=1.5, space_before=0, space_after=0, align=WD_ALIGN_PARAGRAPH.LEFT)
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.font.name = "Arial"
        r_pre.font.size = Pt(11)
        r_pre.font.bold = True
    r = p.add_run(text)
    r.font.name = "Arial"
    r.font.size = Pt(11)
    return p

def add_bullet_p(doc, bullet_label, text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.5)
    p.paragraph_format.first_line_indent = Inches(-0.25)
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    
    r_lbl = p.add_run(bullet_label + " ")
    r_lbl.font.name = "Arial"
    r_lbl.font.size = Pt(11)
    r_lbl.font.bold = True
    
    r = p.add_run(text)
    r.font.name = "Arial"
    r.font.size = Pt(11)
    return p

def add_hanging_ref(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.5)
    p.paragraph_format.first_line_indent = Inches(-0.5)
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(4)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    
    # Process italic tags like *title*
    parts = text.split('*')
    for i, part in enumerate(parts):
        if not part:
            continue
        run = p.add_run(part)
        run.font.name = "Arial"
        run.font.size = Pt(11)
        if i % 2 == 1:
            run.font.italic = True
    return p

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_table_borders(table, color="444444", sz="6"):
    tblPr = table._tbl.tblPr
    borders_xml = f'''
    <w:tblBorders {nsdecls("w")}>
        <w:top w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>
        <w:left w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>
        <w:bottom w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>
        <w:right w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>
        <w:insideH w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>
        <w:insideV w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>
    </w:tblBorders>
    '''
    tblPr.append(parse_xml(borders_xml))

def format_table(table, col_widths):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(table)
    for idx, row in enumerate(table.rows):
        trPr = row._tr.get_or_add_trPr()
        trPr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))
        if idx == 0:
            trPr.append(parse_xml(f'<w:tblHeader {nsdecls("w")}/>'))
        
        for c_idx, cell in enumerate(row.cells):
            cell.width = Inches(col_widths[c_idx])
            set_cell_margins(cell, top=100, bottom=100, left=120, right=120)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.TOP
            if idx == 0:
                # Header row styling
                shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="F2F2F2"/>')
                cell._tc.get_or_add_tcPr().append(shd)
                for p in cell.paragraphs:
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    p.paragraph_format.line_spacing = 1.15
                    p.paragraph_format.space_before = Pt(2)
                    p.paragraph_format.space_after = Pt(2)
                    for r in p.runs:
                        r.font.name = "Arial"
                        r.font.size = Pt(9.5)
                        r.font.bold = True
            else:
                for p in cell.paragraphs:
                    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                    p.paragraph_format.line_spacing = 1.15
                    p.paragraph_format.space_before = Pt(1)
                    p.paragraph_format.space_after = Pt(1)
                    p.paragraph_format.first_line_indent = None
                    for r in p.runs:
                        r.font.name = "Arial"
                        r.font.size = Pt(8.5)

print("Helper functions defined successfully.")
