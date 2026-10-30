"""
audit_typography.py
Detailed inspection of fonts, line spacing, indents, and table properties.
"""

import docx
from docx import Document
from docx.shared import Inches, Pt

DOC_PATH = r"c:\Users\Admin\AndroidStudioProjects\class_realm01\INTEGRATED SENIOR HIGH SCHOOL ATTENDANCE.docx"

doc = Document(DOC_PATH)

print(f"Total paragraphs: {len(doc.paragraphs)}")
print(f"Total tables: {len(doc.tables)}")

# Sample various paragraph types
headings = []
body_samples = []
ref_samples = []

for i, p in enumerate(doc.paragraphs):
    txt = p.text.strip()
    if not txt:
        continue
    
    # Check if heading
    if any(txt.startswith(prefix) for prefix in ["CHAPTER", "1.", "2.", "Table", "Figure"]):
        run_fonts = [r.font.name for r in p.runs if r.font.name]
        run_sizes = [r.font.size.pt for r in p.runs if r.font.size]
        run_bolds = [r.font.bold for r in p.runs if r.font.bold is not None]
        headings.append({
            "idx": i,
            "text": txt[:50],
            "fonts": list(set(run_fonts)),
            "sizes": list(set(run_sizes)),
            "bold": any(run_bolds),
            "space_before": p.paragraph_format.space_before.pt if p.paragraph_format.space_before else None,
            "keep_with_next": p.paragraph_format.keep_with_next
        })
    elif "http" in txt and ("(" in txt or ")" in txt):
        ref_samples.append({
            "idx": i,
            "text": txt[:50],
            "left_indent_in": p.paragraph_format.left_indent.inches if p.paragraph_format.left_indent else None,
            "first_line_indent_in": p.paragraph_format.first_line_indent.inches if p.paragraph_format.first_line_indent else None
        })
    elif i > 30 and len(txt) > 80:
        if len(body_samples) < 5:
            body_samples.append({
                "idx": i,
                "text": txt[:50],
                "first_line_indent_in": p.paragraph_format.first_line_indent.inches if p.paragraph_format.first_line_indent else None,
                "line_spacing": p.paragraph_format.line_spacing,
                "space_before": p.paragraph_format.space_before.pt if p.paragraph_format.space_before else None,
                "space_after": p.paragraph_format.space_after.pt if p.paragraph_format.space_after else None
            })

print("\n--- Heading Samples (Arial 12/11 Bold, keep_with_next) ---")
for h in headings[:8]:
    print(h)

print("\n--- Body Paragraph Samples (0.5\" indent, 1.5 line spacing, 0pt before/after) ---")
for b in body_samples:
    print(b)

print("\n--- Reference Samples (0.5\" hanging indent: left=0.5, firstLine=-0.5) ---")
for r in ref_samples[:3]:
    print(r)

# Table 1 inspection
tbl = doc.tables[0]
print("\n--- Table 1 Formatting Audit ---")
print(f"Alignment: {tbl.alignment}")
header_tr = tbl.rows[0]._tr.get_or_add_trPr()
print(f"Header has tblHeader: {len(header_tr.findall('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}tblHeader')) > 0}")
print(f"Header has cantSplit: {len(header_tr.findall('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}cantSplit')) > 0}")

row1_tr = tbl.rows[1]._tr.get_or_add_trPr()
print(f"Body Row 1 has cantSplit: {len(row1_tr.findall('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}cantSplit')) > 0}")

print("\nTypography and formatting audit complete.")
