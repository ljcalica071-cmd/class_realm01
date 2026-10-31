import os
import zipfile
import xml.etree.ElementTree as ET
from docx import Document
from docx.oxml.ns import qn

def inspect_styles_and_formatting(file_path):
    print("=" * 80)
    print("DEEP FORMATTING ANALYSIS:", file_path)
    print("=" * 80)
    doc = Document(file_path)
    
    # Check paragraphs styles, line spacing, font, indents
    seen_styles = set()
    for p in doc.paragraphs:
        txt = p.text.strip()
        if not txt:
            continue
        s_name = p.style.name if p.style else "NoStyle"
        if s_name not in seen_styles:
            seen_styles.add(s_name)
            p_format = p.paragraph_format
            r_font = p.runs[0].font if p.runs else None
            font_name = r_font.name if r_font else "Inherited"
            font_size = r_font.size.pt if r_font and r_font.size else "Inherited"
            line_spacing = p_format.line_spacing
            first_line_indent = p_format.first_line_indent.inches if p_format.first_line_indent else 0
            space_before = p_format.space_before.pt if p_format.space_before else 0
            space_after = p_format.space_after.pt if p_format.space_after else 0
            align = p_format.alignment
            print(f"Style: {s_name}")
            print(f"  Font: {font_name}, Size: {font_size}")
            print(f"  Line spacing: {line_spacing}, Indent: {first_line_indent} in, Before: {space_before} pt, After: {space_after} pt, Align: {align}")
            print(f"  Sample: {txt[:80]}...\n")
            
    # Sample body paragraph inspection
    print("--- SAMPLE BODY PARAGRAPHS ---")
    body_samples = [p for p in doc.paragraphs if len(p.text.split()) > 20][:5]
    for p in body_samples:
        pf = p.paragraph_format
        r = p.runs[0] if p.runs else None
        print(f"Text ({len(p.text.split())} w): {p.text[:70]}...")
        print(f"  Style: {p.style.name if p.style else 'None'}, Font: {r.font.name if r else 'None'}, Size: {r.font.size.pt if r and r.font.size else 'None'}")
        print(f"  Line spacing: {pf.line_spacing}, Indent: {pf.first_line_indent.inches if pf.first_line_indent else 0} in, Align: {pf.alignment}")
        print(f"  Space Before: {pf.space_before.pt if pf.space_before else 0}, Space After: {pf.space_after.pt if pf.space_after else 0}")
        # check xml properties of paragraph
        pPr = p._p.pPr
        if pPr is not None:
            spacing = pPr.find(qn('w:spacing'))
            if spacing is not None:
                print(f"  w:spacing: line={spacing.get(qn('w:line'))}, lineRule={spacing.get(qn('w:lineRule'))}, before={spacing.get(qn('w:before'))}, after={spacing.get(qn('w:after'))}")
            ind = pPr.find(qn('w:ind'))
            if ind is not None:
                print(f"  w:ind: firstLine={ind.get(qn('w:firstLine'))}, left={ind.get(qn('w:left'))}")

    # Inspect tables
    print("\n--- TABLE XML PROPERTIES ---")
    for t_idx, table in enumerate(doc.tables):
        print(f"Table {t_idx}: {len(table.rows)} rows, {len(table.columns)} cols")
        # Check first row
        if len(table.rows) > 0:
            first_row = table.rows[0]
            trPr = first_row._tr.trPr
            is_tblHeader = trPr.find(qn('w:tblHeader')) is not None if trPr is not None else False
            is_cantSplit = trPr.find(qn('w:cantSplit')) is not None if trPr is not None else False
            print(f"  Row 0: tblHeader={is_tblHeader}, cantSplit={is_cantSplit}")
        # Check second row
        if len(table.rows) > 1:
            second_row = table.rows[1]
            trPr = second_row._tr.trPr
            is_cantSplit = trPr.find(qn('w:cantSplit')) is not None if trPr is not None else False
            print(f"  Row 1: cantSplit={is_cantSplit}")

if __name__ == "__main__":
    base_dir = r"c:\Users\Admin\AndroidStudioProjects\class_realm01"
    fmt_file = os.path.join(base_dir, "Capstone-1-and-Thesis-1-Document-FORMATS-1785742696-1.docx")
    ch123_file = os.path.join(base_dir, "INTEGRATED_SENIOR_HIGH_SCHOOL_ATTENDANCE_CH1_CH2_CH3.docx")
    
    inspect_styles_and_formatting(fmt_file)
    inspect_styles_and_formatting(ch123_file)
