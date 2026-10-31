from docx import Document

doc1 = Document(r"c:\Users\Admin\AndroidStudioProjects\class_realm01\INTEGRATED_SENIOR_HIGH_SCHOOL_ATTENDANCE_CH1_CH2_CH3.docx")
doc2 = Document(r"c:\Users\Admin\AndroidStudioProjects\class_realm01\INTEGRATED SENIOR HIGH SCHOOL ATTENDANCE.docx")

p1 = [p.text.strip() for p in doc1.paragraphs if p.text.strip()]
p2 = [p.text.strip() for p in doc2.paragraphs if p.text.strip()]

print(f"Doc1 non-empty paras: {len(p1)}, Doc2 non-empty paras: {len(p2)}")

diffs = []
for i in range(min(len(p1), len(p2))):
    if p1[i] != p2[i]:
        diffs.append((i, p1[i][:60], p2[i][:60]))

print(f"Differences found: {len(diffs)}")
for d in diffs[:10]:
    print(f"Diff at index {d[0]}:")
    print(f"  Doc1: {d[1]}")
    print(f"  Doc2: {d[2]}")

if len(p1) != len(p2):
    print(f"Extra paragraphs in longer doc: {abs(len(p1)-len(p2))}")
