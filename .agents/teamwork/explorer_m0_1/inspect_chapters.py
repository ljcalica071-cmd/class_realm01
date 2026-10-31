import os
import docx
from docx import Document
from docx.oxml.ns import qn

doc = Document(r"c:\Users\Admin\AndroidStudioProjects\class_realm01\INTEGRATED_SENIOR_HIGH_SCHOOL_ATTENDANCE_CH1_CH2_CH3.docx")

print(f"Total paragraphs: {len(doc.paragraphs)}")
print(f"Total tables: {len(doc.tables)}")

# Print all paragraphs up to Chapter II
print("=" * 80)
print("CHAPTER 1 AND PREAMBLE PARAGRAPHS")
print("=" * 80)
in_ch1 = True
for idx, p in enumerate(doc.paragraphs):
    txt = p.text.strip()
    # Check for page breaks in runs
    has_break = False
    for r in p.runs:
        if 'lastRenderedPageBreak' in r._r.xml or 'w:br' in r._r.xml:
            # check type of br
            for br in r._r.findall(qn('w:br')):
                if br.get(qn('w:type')) == 'page':
                    has_break = True
    br_str = " [PAGE BREAK]" if has_break else ""
    if "CHAPTER II" in txt:
        print(f"--- REACHED CHAPTER II AT P {idx} ---")
        in_ch1 = False
        break
    if txt:
        print(f"P{idx}{br_str} [{p.style.name}]: {txt}")
    elif has_break:
        print(f"P{idx} [PAGE BREAK EMPTY PARAGRAPH]")
