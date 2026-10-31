import os
from docx import Document

doc = Document(r"c:\Users\Admin\AndroidStudioProjects\class_realm01\INTEGRATED_SENIOR_HIGH_SCHOOL_ATTENDANCE_CH1_CH2_CH3.docx")

with open(r"c:\Users\Admin\AndroidStudioProjects\class_realm01\.agents\teamwork\explorer_m0_1\inspect_ch2_ch3_utf8.txt", "w", encoding="utf-8") as out:
    out.write("=" * 80 + "\n")
    out.write("ALL TABLES IN DOCUMENT:\n")
    out.write("=" * 80 + "\n")
    for t_idx, table in enumerate(doc.tables):
        if len(table.rows) > 0:
            c0 = table.rows[0].cells[0].text.strip()
            out.write(f"\n--- Table {t_idx}: {len(table.rows)} rows, {len(table.columns)} cols. First cell: {c0[:50]} ---\n")
            for r_idx, row in enumerate(table.rows):
                vals = [c.text.strip().replace('\n', ' ') for c in row.cells]
                out.write(f"Row {r_idx}: {' | '.join(vals)}\n")

    out.write("\n" + "=" * 80 + "\n")
    out.write("CHAPTER 2 PARAGRAPHS:\n")
    out.write("=" * 80 + "\n")
    in_ch2 = False
    in_ch3 = False

    for idx, p in enumerate(doc.paragraphs):
        txt = p.text.strip()
        if "CHAPTER II" in txt:
            in_ch2 = True
        elif "CHAPTER III" in txt:
            in_ch2 = False
            in_ch3 = True
            out.write("\n" + "=" * 80 + "\n")
            out.write("CHAPTER 3 PARAGRAPHS:\n")
            out.write("=" * 80 + "\n")
        
        if in_ch2:
            if txt:
                out.write(f"[CH2 P{idx}][{len(txt.split())}w] {txt}\n")
        elif in_ch3:
            if txt:
                out.write(f"[CH3 P{idx}][{len(txt.split())}w] {txt}\n")

print("Done writing inspect_ch2_ch3_utf8.txt")
