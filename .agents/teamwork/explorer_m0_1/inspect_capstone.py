import os
import zipfile
import xml.etree.ElementTree as ET
from docx import Document
from docx.shared import Inches, Pt

def inspect_formatting_doc(file_path):
    print("=" * 80)
    print("INSPECTING FORMATTING GUIDELINES DOC:", file_path)
    print("=" * 80)
    doc = Document(file_path)
    
    # Sections and margins
    for i, s in enumerate(doc.sections):
        print(f"Section {i}:")
        print(f"  Page Width: {s.page_width.inches:.2f} in, Height: {s.page_height.inches:.2f} in")
        print(f"  Margins: Top={s.top_margin.inches:.2f} in, Bottom={s.bottom_margin.inches:.2f} in, Left={s.left_margin.inches:.2f} in, Right={s.right_margin.inches:.2f} in")
        print(f"  Header={s.header_distance.inches:.2f} in, Footer={s.footer_distance.inches:.2f} in")

    # Document text
    print("\n--- FORMATTING DOC FULL TEXT ---")
    for p_idx, p in enumerate(doc.paragraphs):
        txt = p.text.strip()
        if txt:
            style_name = p.style.name if p.style else "NoStyle"
            font_name = None
            font_size = None
            if p.runs:
                font_name = p.runs[0].font.name
                font_size = p.runs[0].font.size.pt if p.runs[0].font.size else None
            print(f"[{p_idx}][{style_name}][{font_name} {font_size}pt] {txt}")

    # Inspect tables in formatting doc
    print("\n--- FORMATTING DOC TABLES ---")
    for t_idx, table in enumerate(doc.tables):
        print(f"Table {t_idx} ({len(table.rows)} rows, {len(table.columns)} cols):")
        for r_idx, row in enumerate(table.rows):
            row_txt = [c.text.strip().replace('\n', ' ') for c in row.cells]
            print(f"  Row {r_idx}: {' | '.join(row_txt)}")

def inspect_capstone_doc(file_path, label):
    print("=" * 80)
    print(f"INSPECTING CAPSTONE DOC ({label}): {file_path}")
    print("=" * 80)
    doc = Document(file_path)
    
    total_words = 0
    total_paras = len(doc.paragraphs)
    total_tables = len(doc.tables)
    
    # Check sections / margins
    for i, s in enumerate(doc.sections):
        print(f"Section {i}: Margins Top={s.top_margin.inches:.2f}, Bottom={s.bottom_margin.inches:.2f}, Left={s.left_margin.inches:.2f}, Right={s.right_margin.inches:.2f}")

    # Inspect paragraphs
    headings = []
    current_heading = "Preamble / Title"
    section_word_counts = {current_heading: 0}
    section_para_counts = {current_heading: 0}
    
    for p_idx, p in enumerate(doc.paragraphs):
        words = len(p.text.strip().split())
        total_words += words
        
        # Check if heading
        is_heading = False
        text_strip = p.text.strip()
        style_name = p.style.name if p.style else ""
        
        if "Heading" in style_name or text_strip.startswith("CHAPTER") or (len(text_strip) > 0 and text_strip[0].isdigit() and "." in text_strip[:6] and len(text_strip) < 100):
            current_heading = f"[{p_idx}] {style_name}: {text_strip}"
            headings.append(current_heading)
            section_word_counts[current_heading] = 0
            section_para_counts[current_heading] = 0
        
        section_word_counts[current_heading] += words
        section_para_counts[current_heading] += 1

    print(f"Total Paragraphs: {total_paras}")
    print(f"Total Words (body paragraphs): {total_words}")
    print(f"Total Tables: {total_tables}")
    
    # Check docProps/app.xml for pages
    try:
        with zipfile.ZipFile(file_path, 'r') as z:
            if 'docProps/app.xml' in z.namelist():
                app_xml = z.read('docProps/app.xml')
                root = ET.fromstring(app_xml)
                pages = root.find('{http://schemas.openxmlformats.org/officeDocument/2006/extended-properties}Pages')
                words_prop = root.find('{http://schemas.openxmlformats.org/officeDocument/2006/extended-properties}Words')
                paras_prop = root.find('{http://schemas.openxmlformats.org/officeDocument/2006/extended-properties}Paragraphs')
                lines_prop = root.find('{http://schemas.openxmlformats.org/officeDocument/2006/extended-properties}Lines')
                print(f"App.xml metadata -> Pages: {pages.text if pages is not None else 'N/A'}, Words: {words_prop.text if words_prop is not None else 'N/A'}, Paras: {paras_prop.text if paras_prop is not None else 'N/A'}")
    except Exception as e:
        print(f"Could not read app.xml: {e}")

    print("\n--- OUTLINE & WORD COUNTS PER SECTION ---")
    for h in headings:
        print(f"{h} -> {section_word_counts.get(h, 0)} words, {section_para_counts.get(h, 0)} paras")

if __name__ == "__main__":
    base_dir = r"c:\Users\Admin\AndroidStudioProjects\class_realm01"
    fmt_file = os.path.join(base_dir, "Capstone-1-and-Thesis-1-Document-FORMATS-1785742696-1.docx")
    ch123_file = os.path.join(base_dir, "INTEGRATED_SENIOR_HIGH_SCHOOL_ATTENDANCE_CH1_CH2_CH3.docx")
    att_file = os.path.join(base_dir, "INTEGRATED SENIOR HIGH SCHOOL ATTENDANCE.docx")
    
    inspect_formatting_doc(fmt_file)
    inspect_capstone_doc(ch123_file, "CH1_CH2_CH3")
    inspect_capstone_doc(att_file, "ATTENDANCE")
