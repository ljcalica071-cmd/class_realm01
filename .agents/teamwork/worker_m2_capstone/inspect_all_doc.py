import os
from docx import Document

doc = Document(r"c:\Users\Admin\AndroidStudioProjects\class_realm01\INTEGRATED_SENIOR_HIGH_SCHOOL_ATTENDANCE_CH1_CH2_CH3.docx")

print(f"Total paragraphs: {len(doc.paragraphs)}")
print(f"Total tables: {len(doc.tables)}")

for idx, p in enumerate(doc.paragraphs):
    txt = p.text.strip()
    if txt.startswith("CHAPTER") or txt.startswith("1.") or txt.startswith("2.") or txt.startswith("3.") or "Figure" in txt or "Table" in txt:
        print(f"P{idx}: {txt[:80]}")

print("\n--- TABLES ---")
for t_idx, t in enumerate(doc.tables):
    first_cell = t.rows[0].cells[0].text.strip() if len(t.rows) > 0 and len(t.rows[0].cells) > 0 else "EMPTY"
    print(f"Table {t_idx}: {len(t.rows)} rows, {len(t.columns)} cols. First cell: {first_cell[:60]}")
