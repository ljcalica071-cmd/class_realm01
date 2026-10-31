import docx
import os
import sys
import zipfile
import re

DOCX_PATH = r"c:\Users\Admin\AndroidStudioProjects\class_realm01\INTEGRATED SENIOR HIGH SCHOOL ATTENDANCE.docx"

def run_forensic_inspection():
    print("=== FORENSIC DELIVERABLE INSPECTION ===")
    if not os.path.exists(DOCX_PATH):
        print(f"FATAL: File {DOCX_PATH} does not exist!")
        sys.exit(1)
        
    file_size = os.path.getsize(DOCX_PATH)
    print(f"File Size: {file_size:,} bytes")
    
    doc = docx.Document(DOCX_PATH)
    
    # Paragraphs
    paragraphs = doc.paragraphs
    print(f"Total Paragraphs: {len(paragraphs)}")
    
    # Word count and character count
    full_text = []
    total_words = 0
    total_chars = 0
    for i, p in enumerate(paragraphs):
        t = p.text.strip()
        full_text.append(t)
        words = len(t.split())
        total_words += words
        total_chars += len(t)
        
    print(f"Total Words in Paragraphs: {total_words:,}")
    print(f"Total Characters in Paragraphs: {total_chars:,}")
    
    # Tables
    print(f"Total Tables: {len(doc.tables)}")
    table_words = 0
    for t_idx, table in enumerate(doc.tables):
        rows = len(table.rows)
        cols = len(table.columns) if rows > 0 else 0
        print(f"Table {t_idx+1}: {rows} rows x {cols} cols")
        for row in table.rows:
            for cell in row.cells:
                t = cell.text.strip()
                table_words += len(t.split())
                total_chars += len(t)
                
    print(f"Total Words in Tables: {table_words:,}")
    print(f"Grand Total Words: {total_words + table_words:,}")
    
    # Media assets in zip package
    with zipfile.ZipFile(DOCX_PATH, 'r') as z:
        media_files = [f for f in z.namelist() if f.startswith('word/media/')]
        print(f"Media Assets ({len(media_files)}): {media_files}")
        
    print("\n=== HEADINGS FOUND ===")
    for i, p in enumerate(paragraphs):
        if p.style and 'Heading' in p.style.name:
            print(f"[{p.style.name}] Para {i}: {p.text}")
        elif p.text.isupper() and len(p.text) < 100 and len(p.text) > 3:
            print(f"[ALL CAPS] Para {i}: {p.text}")
            
    print("\n=== COMPLETE EXTRACT TO TEXT DUMP ===")
    dump_path = r"c:\Users\Admin\AndroidStudioProjects\class_realm01\.agents\teamwork\auditor_ch1_ch2_1\extracted_manuscript.txt"
    with open(dump_path, "w", encoding="utf-8") as out:
        out.write(f"=== MANUSCRIPT DUMP (Size: {file_size} bytes) ===\n\n")
        out.write("--- PARAGRAPHS ---\n")
        for i, p in enumerate(paragraphs):
            out.write(f"[{i:03d} | style={p.style.name if p.style else 'None'}]: {p.text}\n")
        out.write("\n--- TABLES ---\n")
        for t_idx, table in enumerate(doc.tables):
            out.write(f"\nTable {t_idx+1}:\n")
            for r_idx, row in enumerate(table.rows):
                row_cells = [c.text.replace('\n', ' ') for c in row.cells]
                out.write(f"Row {r_idx}: " + " | ".join(row_cells) + "\n")
                
    print(f"Extracted manuscript written to: {dump_path}")

if __name__ == '__main__':
    run_forensic_inspection()
