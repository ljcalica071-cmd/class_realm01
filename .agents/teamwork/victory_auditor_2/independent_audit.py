import os
import sys
import docx
from docx.shared import Inches, Pt
from docx.oxml import parse_xml
import xml.etree.ElementTree as ET
import zipfile

DOC_PATH = r"c:\Users\Admin\AndroidStudioProjects\class_realm01\INTEGRATED SENIOR HIGH SCHOOL ATTENDANCE.docx"

def run_audit():
    print(f"Auditing file: {DOC_PATH}")
    if not os.path.exists(DOC_PATH):
        print("ERROR: File does not exist!")
        return
    
    file_size = os.path.getsize(DOC_PATH)
    print(f"File size: {file_size} bytes")

    doc = docx.Document(DOC_PATH)
    
    # 1. Inspect Sections & Margins
    print("\n=== SECTION & MARGIN INSPECTION ===")
    for i, section in enumerate(doc.sections):
        top_in = round(section.top_margin.inches, 2)
        bottom_in = round(section.bottom_margin.inches, 2)
        left_in = round(section.left_margin.inches, 2)
        right_in = round(section.right_margin.inches, 2)
        print(f"Section {i}: Top={top_in}in, Bottom={bottom_in}in, Left={left_in}in, Right={right_in}in")
        print(f"Twips: top={section.top_margin.twips}, bottom={section.bottom_margin.twips}, left={section.left_margin.twips}, right={section.right_margin.twips}")

    # 2. Inspect Word Count & Paragraphs
    print("\n=== CONTENT METRICS ===")
    total_paras = len(doc.paragraphs)
    total_tables = len(doc.tables)
    para_text = [p.text for p in doc.paragraphs]
    full_text = "\n".join(para_text)
    words = full_text.split()
    print(f"Total Paragraphs: {total_paras}")
    print(f"Total Tables: {total_tables}")
    print(f"Total Paragraph Words: {len(words)}")
    print(f"Total Paragraph Characters: {len(full_text)}")

    # 3. Typography Inspection
    print("\n=== TYPOGRAPHY INSPECTION ===")
    font_names = set()
    font_sizes = set()
    table_font_names = set()
    table_font_sizes = set()

    for p in doc.paragraphs:
        for r in p.runs:
            if r.font.name:
                font_names.add(r.font.name)
            if r.font.size:
                font_sizes.add(round(r.font.size.pt, 1))

    for t in doc.tables:
        for row in t.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    for r in p.runs:
                        if r.font.name:
                            table_font_names.add(r.font.name)
                        if r.font.size:
                            table_font_sizes.add(round(r.font.size.pt, 1))

    print(f"Body Font Names detected: {font_names}")
    print(f"Body Font Sizes (pt) detected: {sorted(list(font_sizes))}")
    print(f"Table Font Names detected: {table_font_names}")
    print(f"Table Font Sizes (pt) detected: {sorted(list(table_font_sizes))}")

    # 4. Embedded Media in Zip
    print("\n=== EMBEDDED MEDIA ASSETS ===")
    with zipfile.ZipFile(DOC_PATH, 'r') as z:
        media_files = [f for f in z.namelist() if f.startswith('word/media/')]
        for mf in sorted(media_files):
            info = z.getinfo(mf)
            print(f"Media file: {mf} (size: {info.file_size} bytes)")

    # 5. Table Inspection
    print("\n=== TABLE 1 SYNTHESIS MATRIX INSPECTION ===")
    for i, t in enumerate(doc.tables):
        print(f"Table {i+1}: {len(t.rows)} rows, {len(t.columns)} columns")
        headers = [c.text.strip().replace('\n', ' ') for c in t.rows[0].cells]
        print(f"Header: {headers}")
        for r_idx in range(1, len(t.rows)):
            row_cells = [c.text.strip().replace('\n', ' ') for c in t.rows[r_idx].cells]
            print(f"Row {r_idx}: Author='{row_cells[0][:40]}' | Title='{row_cells[1][:50]}' | Design='{row_cells[2][:30]}' | N='{row_cells[3][:25]}'")

    # Save full text to file for forensic inspection
    output_txt = r"c:\Users\Admin\AndroidStudioProjects\class_realm01\.agents\teamwork\victory_auditor_2\extracted_text.txt"
    with open(output_txt, "w", encoding="utf-8") as f:
        f.write(full_text)
    print(f"\nExtracted full text saved to: {output_txt}")

if __name__ == "__main__":
    run_audit()
