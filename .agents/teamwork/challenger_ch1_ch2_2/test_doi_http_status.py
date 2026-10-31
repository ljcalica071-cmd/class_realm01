"""
test_doi_http_status.py
Dedicated network validation script to test HTTP resolution and validity of all DOIs and URLs
in INTEGRATED SENIOR HIGH SCHOOL ATTENDANCE.docx.
"""

import os
import sys
import re
import urllib.request
import urllib.error
import docx

DOCX_PATH = r"c:\Users\Admin\AndroidStudioProjects\class_realm01\INTEGRATED SENIOR HIGH SCHOOL ATTENDANCE.docx"

def check_url(url, timeout=12):
    req = urllib.request.Request(
        url,
        headers={
            "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/120.0.0.0 Safari/537.36"
        }
    )
    try:
        with urllib.request.urlopen(req, timeout=timeout) as response:
            return response.getcode(), response.geturl(), "OK"
    except urllib.error.HTTPError as e:
        return e.code, url, f"HTTPError: {e.reason}"
    except urllib.error.URLError as e:
        return 0, url, f"URLError: {e.reason}"
    except Exception as e:
        return -1, url, f"Exception: {str(e)}"

def main():
    print("=" * 80)
    print("EMPIRICAL DOI & ACADEMIC URL NETWORK RESOLUTION TEST")
    print("Target Document:", DOCX_PATH)
    print("=" * 80)

    if not os.path.exists(DOCX_PATH):
        print(f"FAIL: Deliverable not found at {DOCX_PATH}")
        sys.exit(1)

    doc = docx.Document(DOCX_PATH)
    full_text = "\n".join([p.text for p in doc.paragraphs])
    
    # Also get text from tables
    table_text = "\n".join(["\n".join([c.text for c in row.cells]) for t in doc.tables for row in t.rows])
    combined_text = full_text + "\n" + table_text

    raw_urls = re.findall(r'https?://[^\s<>"\)\]]+', combined_text)
    # Clean trailing punctuation like periods or commas
    cleaned_urls = []
    for u in raw_urls:
        u_clean = u.rstrip(".,;:)")
        cleaned_urls.append(u_clean)
        
    unique_urls = sorted(list(set(cleaned_urls)))
    print(f"Found {len(unique_urls)} unique URLs to test.\n")

    results = []
    for idx, u in enumerate(unique_urls, 1):
        print(f"[{idx}/{len(unique_urls)}] Testing: {u} ... ", end="", flush=True)
        code, final_url, msg = check_url(u)
        print(f"Code: {code} ({msg})")
        if final_url != u and code != 0:
            print(f"      Redirected to: {final_url}")
        results.append({
            "url": u,
            "code": code,
            "final_url": final_url,
            "status": msg
        })

    print("\n" + "=" * 80)
    print("DOI RESOLUTION SUMMARY")
    print("=" * 80)
    for r in results:
        code = r["code"]
        verdict = "PASS (Active)" if code in [200, 301, 302, 303, 307, 308] else ("PASS (Cloudflare/Paywall 403)" if code == 403 else "FAIL")
        print(f"[{verdict}] HTTP {code:3d} | {r['url']}")
        if r['final_url'] != r['url'] and code != 0:
            print(f"       -> {r['final_url']}")

if __name__ == "__main__":
    main()
