"""
test_deep_analysis.py
Deep forensic inspection of:
- 10 SMART criteria for SY 2026-2027
- 9 Operational Definitions
- Scope boundaries and exclusions
- DeLone & McLean (2003) 6 dimensions
- 4-part IPSO narrative
- Exact sample sizes and methodologies for all 10 studies
"""

import os
import re
import docx

DOCX_PATH = r"c:\Users\Admin\AndroidStudioProjects\class_realm01\INTEGRATED SENIOR HIGH SCHOOL ATTENDANCE.docx"

def main():
    doc = docx.Document(DOCX_PATH)
    paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]

    print("=" * 80)
    print("DEEP FORENSIC ANALYSIS: CHAPTER 1 & CHAPTER 2")
    print("=" * 80)

    # 1. SMART CRITERIA INSPECTION
    print("\n--- 1. SPECIFIC OBJECTIVES & SMART TARGETS ---")
    smart_paras = []
    in_obj = False
    for p in paragraphs:
        if "Specific Objectives" in p:
            in_obj = True
            continue
        if in_obj and ("Scope and Delimitation" in p or "Scope and Limitations" in p or "Significance of the Study" in p):
            in_obj = False
            break
        if in_obj:
            smart_paras.append(p)

    print(f"Total paragraphs in Objectives section: {len(smart_paras)}")
    for idx, sp in enumerate(smart_paras, 1):
        print(f"\n[Objective {idx}]:\n{sp}")

    # 2. OPERATIONAL DEFINITIONS INSPECTION
    print("\n--- 2. DEFINITION OF TERMS INSPECTION ---")
    def_paras = []
    in_def = False
    for p in paragraphs:
        if "Definition of Terms" in p:
            in_def = True
            continue
        if in_def and ("CHAPTER II" in p or "Chapter 2" in p):
            in_def = False
            break
        if in_def:
            def_paras.append(p)

    print(f"Total definitions found: {len(def_paras)}")
    for idx, dp in enumerate(def_paras, 1):
        print(f"\n[Def {idx}]:\n{dp}")

    # 3. SCOPE AND LIMITATIONS INSPECTION
    print("\n--- 3. SCOPE AND LIMITATIONS INSPECTION ---")
    scope_paras = []
    in_scope = False
    for p in paragraphs:
        if "Scope and Delimitation" in p or "Scope and Limitations" in p:
            in_scope = True
            continue
        if in_scope and ("Significance of the Study" in p or "Definition of Terms" in p):
            in_scope = False
            break
        if in_scope:
            scope_paras.append(p)

    print(f"Total paragraphs in Scope: {len(scope_paras)}")
    for idx, sc in enumerate(scope_paras, 1):
        print(f"\n[Scope {idx}]:\n{sc}")

    # 4. TABLE 1 COMPLETE ROWS INSPECTION
    print("\n--- 4. TABLE 1 SYNTHESIS MATRIX ALL 10 ROWS ---")
    if doc.tables:
        t = doc.tables[0]
        for r_idx, row in enumerate(t.rows):
            cells = [c.text.strip().replace("\n", " ") for c in row.cells]
            print(f"\nRow {r_idx} (Len {len(cells)} cols):")
            for c_idx, cell in enumerate(cells):
                print(f"  Col {c_idx+1}: {cell}")

    # 5. THEORETICAL & CONCEPTUAL FRAMEWORK NARRATIVE
    print("\n--- 5. THEORETICAL & CONCEPTUAL FRAMEWORK NARRATIVE ---")
    tf_paras = []
    in_tf = False
    for p in paragraphs:
        if "Theoretical Framework" in p:
            in_tf = True
            continue
        if in_tf and "Conceptual Framework" in p:
            in_tf = False
            print(f"Theoretical Framework Paragraphs: {len(tf_paras)}")
            for idx, tp in enumerate(tf_paras, 1):
                print(f"\n[TF {idx}]:\n{tp}")
            break
        if in_tf:
            tf_paras.append(p)

    cf_paras = []
    in_cf = False
    for p in paragraphs:
        if "Conceptual Framework" in p:
            in_cf = True
            continue
        if in_cf and ("Synthesis" in p or "Chapter II" in p or "References" in p or "Definition of Terms" in p):
            in_cf = False
            print(f"\nConceptual Framework Paragraphs: {len(cf_paras)}")
            for idx, cp in enumerate(cf_paras, 1):
                print(f"\n[CF {idx}]:\n{cp}")
            break
        if in_cf:
            cf_paras.append(p)

if __name__ == "__main__":
    main()
