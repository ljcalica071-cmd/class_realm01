import os
import zipfile
import xml.etree.ElementTree as ET
from docx import Document
from docx.shared import Inches, Pt, Twips
from docx.enum.text import WD_ALIGN_PARAGRAPH

TARGET_DOCX = r"c:\Users\Admin\AndroidStudioProjects\class_realm01\INTEGRATED SENIOR HIGH SCHOOL ATTENDANCE.docx"
BACKUP_DOCX = r"c:\Users\Admin\AndroidStudioProjects\class_realm01\.agents\teamwork\worker_ch1_ch2_compiler\INTEGRATED SENIOR HIGH SCHOOL ATTENDANCE.backup.docx"

W_NAMESPACE = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
R_NAMESPACE = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
A_NAMESPACE = "http://schemas.openxmlformats.org/drawingml/2006/main"
WP_NAMESPACE = "http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing"
NS = {
    'w': W_NAMESPACE,
    'r': R_NAMESPACE,
    'a': A_NAMESPACE,
    'wp': WP_NAMESPACE
}

def register_namespaces():
    for prefix, uri in NS.items():
        ET.register_namespace(prefix, uri)

register_namespaces()

def run_all_tests():
    print("=" * 80)
    print("CHALLENGER 1: INDEPENDENT EMPIRICAL VERIFICATION & STRESS TEST")
    print(f"Target: {TARGET_DOCX}")
    print("=" * 80)
    
    if not os.path.exists(TARGET_DOCX):
        print(f"FAIL: File does not exist: {TARGET_DOCX}")
        return False
        
    file_size = os.path.getsize(TARGET_DOCX)
    print(f"File Size: {file_size:,} bytes")
    
    # Open docx with python-docx and zipfile
    doc = Document(TARGET_DOCX)
    
    with zipfile.ZipFile(TARGET_DOCX, 'r') as zf:
        doc_xml_bytes = zf.read('word/document.xml')
        media_files = [f for f in zf.namelist() if f.startswith('word/media/')]
        header_files = [f for f in zf.namelist() if 'header' in f]
        footer_files = [f for f in zf.namelist() if 'footer' in f]
        rels_xml_bytes = zf.read('word/_rels/document.xml.rels')

    tree = ET.fromstring(doc_xml_bytes)
    body = tree.find('w:body', NS)
    body_elements = list(body)
    print(f"Total top-level w:body child elements: {len(body_elements)}")
    
    all_passed = True
    
    # -------------------------------------------------------------------------
    # TEST 1: Title Page Front Matter Preservation (elements 0 to 27)
    # -------------------------------------------------------------------------
    print("\n" + "=" * 60)
    print("TEST 1: TITLE PAGE FRONT MATTER PRESERVATION (ELEMENTS 0-27)")
    print("=" * 60)
    test1_passed = True
    
    # Check if backup exists
    has_backup = os.path.exists(BACKUP_DOCX)
    if has_backup:
        with zipfile.ZipFile(BACKUP_DOCX, 'r') as bzf:
            backup_doc_xml = bzf.read('word/document.xml')
            backup_tree = ET.fromstring(backup_doc_xml)
            backup_body = backup_tree.find('w:body', NS)
            backup_elements = list(backup_body)
            print(f"Backup baseline total body elements: {len(backup_elements)}")
            
            # Compare elements 0 to 27 byte-by-byte in XML serialization
            print(f"Comparing elements 0..27 between Target and Baseline Backup...")
            mismatch_count = 0
            for idx in range(28):
                target_elem_xml = ET.tostring(body_elements[idx], encoding='utf-8')
                backup_elem_xml = ET.tostring(backup_elements[idx], encoding='utf-8')
                if target_elem_xml != backup_elem_xml:
                    print(f"  [MISMATCH] Element {idx} differs from baseline backup!")
                    mismatch_count += 1
                else:
                    # Quick text preview
                    text_content = "".join(body_elements[idx].itertext()).strip()
                    if text_content:
                        preview = text_content[:60] + "..." if len(text_content) > 60 else text_content
                        print(f"  [MATCH] Element {idx:02d}: '{preview}'")
                    else:
                        print(f"  [MATCH] Element {idx:02d}: <{body_elements[idx].tag.split('}')[-1]}> (blank/break)")
            
            if mismatch_count == 0:
                print(">> Element 0-27 exact XML match with backup: 100% IDENTICAL (PASS)")
            else:
                print(f">> Element 0-27 differences found: {mismatch_count} (FAIL)")
                test1_passed = False
    else:
        print("[WARN] Baseline backup docx not found for direct byte comparison.")
        
    # Check specific required contents in elements 0 to 27
    elem0_text = "".join(body_elements[0].itertext()).strip()
    print(f"\nElement 0 text: '{elem0_text}'")
    if "INTEGRATED SENIOR HIGH SCHOOL ATTENDANCE" in elem0_text:
        print("  - Title text verified: PASS")
    else:
        print("  - Title text missing: FAIL")
        test1_passed = False
        
    front_text = "".join("".join(e.itertext()) for e in body_elements[:28])
    if "Laguna University" in front_text:
        print("  - Laguna University text verified: PASS")
    else:
        print("  - Laguna University text missing: FAIL")
        test1_passed = False
        
    if "College of Computing Studies" in front_text:
        print("  - College of Computing Studies verified: PASS")
    else:
        print("  - College of Computing Studies missing: FAIL")
        test1_passed = False
        
    if "Bachelor of Science in Information Technology" in front_text or "BSIT" in front_text:
        print("  - BSIT Degree text verified: PASS")
    else:
        print("  - BSIT Degree text missing: FAIL")
        test1_passed = False
        
    if "Calica, Ljay H." in front_text and "San Juan, John Rex G." in front_text:
        print("  - Author names (Calica, Ljay H. & San Juan, John Rex G.) verified: PASS")
    else:
        print("  - Author names missing or altered: FAIL")
        test1_passed = False
        
    # Check element 27 page break
    elem27 = body_elements[27]
    elem27_xml = ET.tostring(elem27, encoding='utf-8').decode('utf-8')
    has_page_break = '<w:br w:type="page"' in elem27_xml or 'w:type="page"' in elem27_xml
    print(f"Element 27 tag: <{elem27.tag.split('}')[-1]}>, has page break: {has_page_break}")
    print(f"Element 27 XML snippet: {elem27_xml[:150]}")
    if has_page_break:
        print("  - Element 27 explicit page break verified: PASS")
    else:
        print("  - Element 27 explicit page break missing: FAIL")
        test1_passed = False
        
    # Check headers / logos
    print(f"Header files found in docx: {header_files}")
    if header_files:
        with zipfile.ZipFile(TARGET_DOCX, 'r') as zf:
            header1_xml = zf.read(header_files[0]).decode('utf-8', errors='ignore')
            has_drawing_in_header = '<w:drawing' in header1_xml or '<v:shape' in header1_xml or 'r:id' in header1_xml
            print(f"  - Header 1 contains drawing/images: {has_drawing_in_header}")
    
    if test1_passed:
        print(">> TEST 1 RESULT: PASS")
    else:
        print(">> TEST 1 RESULT: FAIL")
        all_passed = False

    # -------------------------------------------------------------------------
    # TEST 2: Page Margins Compliance
    # -------------------------------------------------------------------------
    print("\n" + "=" * 60)
    print("TEST 2: PAGE MARGINS COMPLIANCE")
    print("=" * 60)
    test2_passed = True
    
    # Locate all sectPr elements in document.xml
    sect_elements = body.findall('.//w:sectPr', NS)
    if not sect_elements:
        # Check last element
        if body_elements[-1].tag.endswith('sectPr'):
            sect_elements = [body_elements[-1]]
            
    print(f"Found {len(sect_elements)} w:sectPr element(s).")
    for s_idx, sect in enumerate(sect_elements):
        pgMar = sect.find('w:pgMar', NS)
        if pgMar is not None:
            top = pgMar.get(f'{{{W_NAMESPACE}}}top')
            bottom = pgMar.get(f'{{{W_NAMESPACE}}}bottom')
            left = pgMar.get(f'{{{W_NAMESPACE}}}left')
            right = pgMar.get(f'{{{W_NAMESPACE}}}right')
            
            top_twips = int(top) if top else 0
            bottom_twips = int(bottom) if bottom else 0
            left_twips = int(left) if left else 0
            right_twips = int(right) if right else 0
            
            top_in = top_twips / 1440.0
            bottom_in = bottom_twips / 1440.0
            left_in = left_twips / 1440.0
            right_in = right_twips / 1440.0
            
            print(f"Section {s_idx} Margins:")
            print(f"  Top:    {top_twips} twips ({top_in:.2f} in) [Expected: 2301 twips / 1.60 in]")
            print(f"  Left:   {left_twips} twips ({left_in:.2f} in) [Expected: 2301 twips / 1.60 in]")
            print(f"  Bottom: {bottom_twips} twips ({bottom_in:.2f} in) [Expected: 1729 twips / 1.20 in]")
            print(f"  Right:  {right_twips} twips ({right_in:.2f} in) [Expected: 2017 twips / 1.40 in]")
            
            if top_twips == 2301 and left_twips == 2301 and bottom_twips == 1729 and right_twips == 2017:
                print("  - Exact twips match institutional requirement: PASS")
            else:
                print("  - Margins deviate from institutional requirement: FAIL")
                test2_passed = False
        else:
            print(f"Section {s_idx} has no w:pgMar element! FAIL")
            test2_passed = False
            
    if test2_passed:
        print(">> TEST 2 RESULT: PASS")
    else:
        print(">> TEST 2 RESULT: FAIL")
        all_passed = False

    # -------------------------------------------------------------------------
    # TEST 3: Typography & Spacing Properties
    # -------------------------------------------------------------------------
    print("\n" + "=" * 60)
    print("TEST 3: TYPOGRAPHY & SPACING PROPERTIES")
    print("=" * 60)
    test3_passed = True
    
    # Inspect paragraphs in docx
    headings_checked = 0
    body_checked = 0
    heading_issues = []
    body_issues = []
    
    for p_idx, p in enumerate(doc.paragraphs):
        p_text = p.text.strip()
        if not p_text:
            continue
            
        # Check if heading
        p_elem = p._element
        pPr = p_elem.pPr
        
        # Determine if heading
        is_heading1 = False
        is_heading2 = False
        if p.style.name.startswith('Heading 1') or p_text.startswith('CHAPTER ') or p_text in [
            'INTRODUCTION', 'REVIEW OF RELATED LITERATURE AND SYSTEMS', 'REFERENCES'
        ]:
            is_heading1 = True
        elif p.style.name.startswith('Heading 2') or any(p_text.startswith(x) for x in [
            'Background of the Study', 'Statement of the Problem', 'Objectives of the Study',
            'General Objective', 'Specific Objectives', 'Significance of the Study',
            'Scope and Limitations', 'Definition of Terms', 'Review of Related Literature',
            'Synthesis of the Reviewed Literature', 'Synthesis Matrix', 'Theoretical Framework',
            'Conceptual Framework'
        ]):
            is_heading2 = True
            
        if is_heading1 or is_heading2:
            headings_checked += 1
            h_type = "Heading 1" if is_heading1 else "Heading 2"
            
            # Check keep_with_next
            has_keep_next = pPr.keepNext is not None if pPr is not None else False
            
            # Check font and bold
            fonts = set()
            bolds = set()
            sizes = set()
            for r in p.runs:
                if r.text.strip():
                    fonts.add(r.font.name or (p.style.font.name if p.style else None))
                    bolds.add(r.bold if r.bold is not None else (p.style.font.bold if p.style else None))
                    sizes.add(r.font.size.pt if r.font.size else (p.style.font.size.pt if p.style and p.style.font.size else None))
            
            # Check spacing before
            space_before_pt = p.paragraph_format.space_before.pt if p.paragraph_format.space_before else 0
            
            status = "OK"
            reasons = []
            if not has_keep_next:
                reasons.append("keepNext=False")
            if not any('Arial' in str(f) for f in fonts):
                reasons.append(f"font={fonts}")
            if False in bolds or None in bolds and len(bolds) == 1:
                reasons.append(f"bold={bolds}")
            if space_before_pt < 10:  # Expect ~1 blank line before (~14pt)
                reasons.append(f"space_before={space_before_pt}pt")
                
            if reasons:
                heading_issues.append((p_idx, p_text[:40], h_type, reasons))
            # print sample
            if headings_checked <= 5:
                print(f"  [HEADING] '{p_text[:35]}...' | Type: {h_type} | Fonts: {fonts} | Bold: {bolds} | Sizes: {sizes} | KeepNext: {has_keep_next} | SpaceBefore: {space_before_pt}pt")
                
        elif p_idx > 27 and len(p_text) > 80 and not p_text.startswith("Figure ") and not p_text.startswith("Table "):
            # Likely narrative body paragraph (skip references for body test)
            if p_idx < 155: # Narrative body before References
                body_checked += 1
                
                # Check indent
                first_line_indent = p.paragraph_format.first_line_indent
                indent_twips = first_line_indent.twips if first_line_indent else 0
                
                # Check line spacing
                line_spacing = p.paragraph_format.line_spacing
                
                # Check space before and after
                sp_before = p.paragraph_format.space_before.pt if p.paragraph_format.space_before else 0
                sp_after = p.paragraph_format.space_after.pt if p.paragraph_format.space_after else 0
                
                # Check font
                fonts = set()
                sizes = set()
                bolds = set()
                for r in p.runs:
                    if r.text.strip():
                        fonts.add(r.font.name or (p.style.font.name if p.style else None))
                        sizes.add(r.font.size.pt if r.font.size else (p.style.font.size.pt if p.style and p.style.font.size else None))
                        bolds.add(r.bold)
                        
                reasons = []
                if indent_twips != 720: # 0.5 in
                    reasons.append(f"indent={indent_twips}twips")
                if line_spacing != 1.5:
                    reasons.append(f"line_spacing={line_spacing}")
                if sp_before != 0:
                    reasons.append(f"sp_before={sp_before}")
                if sp_after != 0:
                    reasons.append(f"sp_after={sp_after}")
                if not any('Arial' in str(f) for f in fonts):
                    reasons.append(f"font={fonts}")
                if not any(s == 11.0 for s in sizes):
                    reasons.append(f"sizes={sizes}")
                    
                if reasons:
                    body_issues.append((p_idx, p_text[:40], reasons))
                    
                if body_checked <= 5:
                    print(f"  [BODY] '{p_text[:35]}...' | Indent: {indent_twips} twips | Spacing: {line_spacing} | Before: {sp_before}pt | After: {sp_after}pt | Font: {fonts} | Size: {sizes}")

    print(f"\nTotal Headings Checked: {headings_checked}, Issues: {len(heading_issues)}")
    if heading_issues:
        for idx, txt, ht, reas in heading_issues[:5]:
            print(f"    Issue at P{idx} ('{txt}'): {', '.join(reas)}")
        test3_passed = False
        
    print(f"Total Narrative Body Paragraphs Checked: {body_checked}, Issues: {len(body_issues)}")
    if body_issues:
        for idx, txt, reas in body_issues[:5]:
            print(f"    Issue at P{idx} ('{txt}'): {', '.join(reas)}")
        test3_passed = False

    if test3_passed:
        print(">> TEST 3 RESULT: PASS")
    else:
        print(">> TEST 3 RESULT: FAIL")
        all_passed = False

    # -------------------------------------------------------------------------
    # TEST 4: Table 1 Synthesis Matrix Structure
    # -------------------------------------------------------------------------
    print("\n" + "=" * 60)
    print("TEST 4: TABLE 1 SYNTHESIS MATRIX STRUCTURE")
    print("=" * 60)
    test4_passed = True
    
    if len(doc.tables) == 0:
        print("FAIL: No tables found in document!")
        test4_passed = False
    else:
        t = doc.tables[0]
        rows = t.rows
        cols = t.columns
        print(f"Table 1 Dimensions: {len(rows)} rows x {len(cols)} columns")
        if len(rows) != 11 or len(cols) != 6:
            print(f"  [WARN] Expected 11 rows (1 header + 10 studies) and 6 columns. Got {len(rows)}x{len(cols)}")
        else:
            print("  - Table dimensions (11x6) verified: PASS")
            
        t_elem = t._element
        tblPr = t_elem.tblPr
        
        # Check borders
        tblBorders = tblPr.find('w:tblBorders', NS)
        if tblBorders is not None:
            border_checks = {}
            for b_name in ['top', 'left', 'bottom', 'right', 'insideH', 'insideV']:
                b = tblBorders.find(f'w:{b_name}', NS)
                if b is not None:
                    val = b.get(f'{{{W_NAMESPACE}}}val')
                    sz = b.get(f'{{{W_NAMESPACE}}}sz')
                    color = b.get(f'{{{W_NAMESPACE}}}color')
                    border_checks[b_name] = (val, sz, color)
                    print(f"  Border {b_name}: val={val}, sz={sz} (0.75pt), color=#{color}")
                    if val != 'single' or sz != '6' or color.upper() != '444444':
                        print(f"    [FAIL] Border {b_name} does not match required single sz=6 #444444!")
                        test4_passed = False
                else:
                    print(f"  [FAIL] Border {b_name} missing!")
                    test4_passed = False
        else:
            print("  [FAIL] tblBorders missing!")
            test4_passed = False
            
        # Check row properties: tblHeader on row 0, cantSplit on all rows
        print("\nChecking Table Row Properties:")
        row0 = rows[0]._element
        trPr0 = row0.trPr
        has_tbl_header = trPr0.find('w:tblHeader', NS) is not None if trPr0 is not None else False
        print(f"  Row 0 <w:tblHeader/> present: {has_tbl_header}")
        if not has_tbl_header:
            print("    [FAIL] Header row 0 is missing <w:tblHeader/>!")
            test4_passed = False
            
        all_cant_split = True
        for r_idx, r in enumerate(rows):
            r_elem = r._element
            r_trPr = r_elem.trPr
            has_cant_split = r_trPr.find('w:cantSplit', NS) is not None if r_trPr is not None else False
            if not has_cant_split:
                print(f"    [FAIL] Row {r_idx} missing <w:cantSplit/>!")
                all_cant_split = False
                test4_passed = False
        if all_cant_split:
            print("  - <w:cantSplit/> present on all 11 rows: PASS")
            
        # Check header row text: uppercase bold text and centered alignment
        print("\nChecking Header Row Text & Styling:")
        header_cells = rows[0].cells
        for c_idx, cell in enumerate(header_cells):
            cell_text = cell.text.strip()
            # check uppercase
            is_upper = cell_text == cell_text.upper()
            
            # check bold and alignment in paragraphs of cell
            cell_bold = True
            cell_centered = True
            for cp in cell.paragraphs:
                if cp.alignment != WD_ALIGN_PARAGRAPH.CENTER:
                    cell_centered = False
                for cr in cp.runs:
                    if cr.text.strip() and not cr.bold:
                        cell_bold = False
                        
            print(f"  Cell {c_idx}: '{cell_text}' | Uppercase: {is_upper} | Bold: {cell_bold} | Centered: {cell_centered}")
            if not is_upper or not cell_bold or not cell_centered:
                print(f"    [FAIL] Header cell {c_idx} does not meet uppercase bold centered criteria!")
                test4_passed = False
                
    if test4_passed:
        print(">> TEST 4 RESULT: PASS")
    else:
        print(">> TEST 4 RESULT: FAIL")
        all_passed = False

    # -------------------------------------------------------------------------
    # TEST 5: Embedded Figures & Captions
    # -------------------------------------------------------------------------
    print("\n" + "=" * 60)
    print("TEST 5: EMBEDDED FIGURES & CAPTIONS")
    print("=" * 60)
    test5_passed = True
    
    print(f"Media files inside word/media/: {media_files}")
    if len(media_files) < 2:
        print(f"FAIL: Expected at least 2 media files for Figures 1 & 2. Found {len(media_files)}")
        test5_passed = False
    else:
        with zipfile.ZipFile(TARGET_DOCX, 'r') as zf:
            for mf in media_files:
                info = zf.getinfo(mf)
                print(f"  - {mf}: {info.file_size:,} bytes")
                
    # Search for Figure 1 and Figure 2 captions and image shapes in paragraphs
    fig1_found = False
    fig2_found = False
    fig1_caption_p = None
    fig2_caption_p = None
    fig1_img_p = None
    fig2_img_p = None
    fig1_note_p = None
    fig2_note_p = None
    
    for p_idx, p in enumerate(doc.paragraphs):
        p_text = p.text.strip()
        if p_text.startswith("Figure 1.") or p_text.startswith("Figure 1:"):
            fig1_found = True
            fig1_caption_p = (p_idx, p)
            # Next paragraph might be the image
            if p_idx + 1 < len(doc.paragraphs):
                fig1_img_p = (p_idx + 1, doc.paragraphs[p_idx + 1])
            if p_idx + 2 < len(doc.paragraphs):
                fig1_note_p = (p_idx + 2, doc.paragraphs[p_idx + 2])
        elif p_text.startswith("Figure 2.") or p_text.startswith("Figure 2:"):
            fig2_found = True
            fig2_caption_p = (p_idx, p)
            if p_idx + 1 < len(doc.paragraphs):
                fig2_img_p = (p_idx + 1, doc.paragraphs[p_idx + 1])
            if p_idx + 2 < len(doc.paragraphs):
                fig2_note_p = (p_idx + 2, doc.paragraphs[p_idx + 2])
                
    print(f"\nFigure 1 Caption detected: {fig1_found}")
    if fig1_caption_p:
        idx, p = fig1_caption_p
        print(f"  Paragraph {idx}: '{p.text}'")
        runs = p.runs
        # Label bold, title italic
        print(f"  Runs count: {len(runs)}")
        for r_i, r in enumerate(runs):
            print(f"    Run {r_i}: '{r.text}' | Bold={r.bold} | Italic={r.italic}")
        label_bold = any(r.bold and "Figure 1" in r.text for r in runs)
        title_italic = any(r.italic for r in runs)
        print(f"  Label bold: {label_bold}, Title italic: {title_italic}")
        if not label_bold or not title_italic:
            print("    [FAIL] Figure 1 caption formatting incorrect!")
            test5_passed = False
            
    if fig1_img_p:
        idx, p = fig1_img_p
        has_drawing = '<w:drawing' in p._element.xml
        is_centered = p.alignment == WD_ALIGN_PARAGRAPH.CENTER
        print(f"  Image P{idx}: has drawing XML: {has_drawing}, centered: {is_centered}")
        if not has_drawing or not is_centered:
            print("    [FAIL] Figure 1 image drawing missing or not centered!")
            test5_passed = False
            
    if fig1_note_p:
        idx, p = fig1_note_p
        print(f"  Citation/Note P{idx}: '{p.text[:70]}...'")
        
    print(f"\nFigure 2 Caption detected: {fig2_found}")
    if fig2_caption_p:
        idx, p = fig2_caption_p
        print(f"  Paragraph {idx}: '{p.text}'")
        runs = p.runs
        for r_i, r in enumerate(runs):
            print(f"    Run {r_i}: '{r.text}' | Bold={r.bold} | Italic={r.italic}")
        label_bold = any(r.bold and "Figure 2" in r.text for r in runs)
        title_italic = any(r.italic for r in runs)
        print(f"  Label bold: {label_bold}, Title italic: {title_italic}")
        if not label_bold or not title_italic:
            print("    [FAIL] Figure 2 caption formatting incorrect!")
            test5_passed = False
            
    if fig2_img_p:
        idx, p = fig2_img_p
        has_drawing = '<w:drawing' in p._element.xml
        is_centered = p.alignment == WD_ALIGN_PARAGRAPH.CENTER
        print(f"  Image P{idx}: has drawing XML: {has_drawing}, centered: {is_centered}")
        if not has_drawing or not is_centered:
            print("    [FAIL] Figure 2 image drawing missing or not centered!")
            test5_passed = False
            
    if fig2_note_p:
        idx, p = fig2_note_p
        print(f"  Citation/Note P{idx}: '{p.text[:70]}...'")
        
    if test5_passed and fig1_found and fig2_found:
        print(">> TEST 5 RESULT: PASS")
    else:
        print(">> TEST 5 RESULT: FAIL")
        all_passed = False

    # -------------------------------------------------------------------------
    # TEST 6: References Layout
    # -------------------------------------------------------------------------
    print("\n" + "=" * 60)
    print("TEST 6: REFERENCES LAYOUT")
    print("=" * 60)
    test6_passed = True
    
    # Locate REFERENCES heading
    ref_heading_idx = None
    for p_idx, p in enumerate(doc.paragraphs):
        if p.text.strip() == "REFERENCES":
            ref_heading_idx = p_idx
            break
            
    print(f"REFERENCES heading found at paragraph index: {ref_heading_idx}")
    if ref_heading_idx is None:
        print("FAIL: REFERENCES heading not found!")
        test6_passed = False
    else:
        ref_paragraphs = doc.paragraphs[ref_heading_idx + 1:]
        print(f"Found {len(ref_paragraphs)} reference paragraph entries.")
        if len(ref_paragraphs) < 10:
            print(f"[FAIL] Expected at least 10-14 reference entries. Found {len(ref_paragraphs)}")
            test6_passed = False
            
        ref_issues = []
        for r_idx, rp in enumerate(ref_paragraphs):
            rtxt = rp.text.strip()
            if not rtxt:
                continue
                
            # Check hanging indent: left=720 twips, first_line=-720 twips (or hanging=720)
            p_elem = rp._element
            pPr = p_elem.pPr
            w_ind = pPr.find('w:ind', NS) if pPr is not None else None
            
            left_twips = None
            first_line_twips = None
            hanging_twips = None
            if w_ind is not None:
                left_twips = w_ind.get(f'{{{W_NAMESPACE}}}left')
                first_line_twips = w_ind.get(f'{{{W_NAMESPACE}}}firstLine')
                hanging_twips = w_ind.get(f'{{{W_NAMESPACE}}}hanging')
                
            # Line spacing
            line_spacing = rp.paragraph_format.line_spacing
            
            # Fonts
            fonts = set()
            sizes = set()
            for r in rp.runs:
                if r.text.strip():
                    fonts.add(r.font.name or (rp.style.font.name if rp.style else None))
                    sizes.add(r.font.size.pt if r.font.size else (rp.style.font.size.pt if rp.style and rp.style.font.size else None))
                    
            reasons = []
            # Check hanging indent: left must be 720 and hanging=720 (or firstLine=-720)
            is_hanging_720 = (left_twips == '720' and (hanging_twips == '720' or first_line_twips == '-720'))
            if not is_hanging_720:
                reasons.append(f"ind(left={left_twips}, hanging={hanging_twips}, firstLine={first_line_twips})")
            if line_spacing != 1.5:
                reasons.append(f"line_spacing={line_spacing}")
            if not any('Arial' in str(f) for f in fonts):
                reasons.append(f"font={fonts}")
            if not any(s == 11.0 for s in sizes):
                reasons.append(f"size={sizes}")
                
            if reasons:
                ref_issues.append((r_idx, rtxt[:50], reasons))
                
            if r_idx < 5:
                print(f"  [REF {r_idx}] '{rtxt[:45]}...' | Ind: left={left_twips}, hang={hanging_twips}, firstLine={first_line_twips} | LineSp: {line_spacing} | Font: {fonts} | Size: {sizes}")

        if ref_issues:
            print(f"Found {len(ref_issues)} reference formatting issues:")
            for r_idx, rtxt, reas in ref_issues:
                print(f"    Issue at Ref {r_idx} ('{rtxt}'): {', '.join(reas)}")
            test6_passed = False
        else:
            print(f"  - All {len(ref_paragraphs)} references verified (0.5\" hanging indent, 1.5 line spacing, Arial 11pt): PASS")
            
    if test6_passed:
        print(">> TEST 6 RESULT: PASS")
    else:
        print(">> TEST 6 RESULT: FAIL")
        all_passed = False
        
    print("\n" + "=" * 80)
    if all_passed:
        print("OVERALL VERDICT: ALL EMPIRICAL TESTS PASSED (APPROVE)")
    else:
        print("OVERALL VERDICT: FAILURES DETECTED (REQUEST_CHANGES)")
    print("=" * 80)
    return all_passed

if __name__ == '__main__':
    run_all_tests()
