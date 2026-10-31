from docx import Document
from docx.oxml.ns import qn

TARGET_DOCX = r"c:\Users\Admin\AndroidStudioProjects\class_realm01\INTEGRATED SENIOR HIGH SCHOOL ATTENDANCE.docx"
doc = Document(TARGET_DOCX)

p30 = doc.paragraphs[30] # 1.1 Project Context
p31 = doc.paragraphs[31] # In the contemporary Philippine...

print(f"P30 text: {p30.text[:30]}")
print(f"P30 paragraph_format.keep_with_next: {p30.paragraph_format.keep_with_next}")
print(f"P30 find keepNext: {p30._element.pPr.find(qn('w:keepNext')) is not None}")

print(f"P31 text: {p31.text[:30]}")
print(f"P31 paragraph_format.keep_with_next: {p31.paragraph_format.keep_with_next}")
print(f"P31 find keepNext: {p31._element.pPr.find(qn('w:keepNext')) is not None}")
