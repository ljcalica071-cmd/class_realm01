r"""
CHALLENGER 1 - INDEPENDENT EMPIRICAL AUDIT & VERIFICATION SUITE
ClassRealm Chapters 1 and 2 Physical & OpenXML Deliverable Verification
Target: c:\Users\Admin\AndroidStudioProjects\class_realm01\INTEGRATED SENIOR HIGH SCHOOL ATTENDANCE.docx
"""

import os
import zipfile
import xml.etree.ElementTree as ET
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

TARGET_DOCX = r"c:\Users\Admin\AndroidStudioProjects\class_realm01\INTEGRATED SENIOR HIGH SCHOOL ATTENDANCE.docx"
BACKUP_DOCX = r"c:\Users\Admin\AndroidStudioProjects\class_realm01\.agents\teamwork\worker_ch1_ch2_compiler\INTEGRATED SENIOR HIGH SCHOOL ATTENDANCE.backup.docx"

W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
NS = {'w': W_NS}

for prefix, uri in [('w', W_NS), ('r', "http://schemas.openxmlformats.org/officeDocument/2006/relationships")]:
    ET.register_namespace(prefix, uri)

def run_suite():
    print("=" * 80)
    print("CHALLENGER 1: COMPREHENSIVE EMPIRICAL VERIFICATION HARNESS")
    print("=" * 80)
    
    assert os.path.exists(TARGET_DOCX), f"Deliverable docx missing at {TARGET_DOCX}"
    file_size = os.path.getsize(TARGET_DOCX)
    print(f"Deliverable Path: {TARGET_DOCX}")
    print(f"Deliverable Size: {file_size:,} bytes")
    
    doc = Document(TARGET_DOCX)
    
    with zipfile.ZipFile(TARGET_DOCX, 'r') as zf:
        doc_xml = zf.read('word/document.xml')
        media_files = [f for f in zf.namelist() if f.startswith('word/media/')]
        header_files = [f for f in zf.namelist() if 'header' in f]
        
    tree = ET.fromstring(doc_xml)
    body = tree.find('w:body', NS)
    body_elems = list(body)
    
    results = {}
    
    # -------------------------------------------------------------------------
    # TEST 1: Title Page Front Matter Preservation (Elements 0 to 27)
    # -------------------------------------------------------------------------
    print("\n[TEST 1] Title Page Front Matter Preservation (Elements 0 - 27)")
    t1_pass = True
    
    assert os.path.exists(BACKUP_DOCX), "Backup file missing!"
    with zipfile.ZipFile(BACKUP_DOCX, 'r') as bzf:
        b_tree = ET.fromstring(bzf.read('word/document.xml'))
        b_body = b_tree.find('w:body', NS)
        b_elems = list(b_body)
        
    print(f"Target body elements: {len(body_elems)}, Backup baseline elements: {len(b_elems)}")
    
    # 1.1 Exact XML comparison for elements 0 to 27
    identical_elements = 0
    for i in range(28):
        t_xml = ET.tostring(body_elems[i], encoding='utf-8')
        b_xml = ET.tostring(b_elems[i], encoding='utf-8')
        if t_xml == b_xml:
            identical_elements += 1
        else:
            print(f"  [FAIL] Element {i} differs from baseline backup!")
            t1_pass = False
            
    print(f"Elements 0-27 identical to baseline: {identical_elements} / 28")
    assert identical_elements == 28, "Not all 28 front matter elements are identical to baseline!"
    
    # 1.2 Specific content checks in elements 0 to 27
    e0_text = "".join(body_elems[0].itertext()).strip()
    assert "INTEGRATED SENIOR HIGH SCHOOL ATTENDANCE" in e0_text
    print(f"  - Element 0 Title: '{e0_text[:70]}...' -> OK")
    
    front_text = " ".join("".join(e.itertext()).strip() for e in body_elems[:28])
    assert "Laguna University" in front_text
    assert "College of Computing Studies" in front_text
    print("  - Institutional Headers (Laguna University, CCS) -> OK")
    
    assert "Bachelor of Science in" in front_text and "Information Technology" in front_text
    print("  - Degree Name (Bachelor of Science in Information Technology) -> OK")
    
    assert "Calica, Ljay H." in front_text and "San Juan, John Rex G." in front_text
    print("  - Authors (Calica, Ljay H. & San Juan, John Rex G.) -> OK")
    
    # Element 27 explicit page break
    br_elem = body_elems[27].find(f'.//{{{W_NS}}}br')
    assert br_elem is not None, "Element 27 has no w:br tag!"
    assert br_elem.get(f'{{{W_NS}}}type') == 'page', f"Element 27 br is not page type: {br_elem.attrib}"
    print("  - Element 27 Explicit Page Break (<w:br w:type=\"page\"/>) -> OK")
    
    # Header logos in word/header1.xml
    with zipfile.ZipFile(TARGET_DOCX, 'r') as zf:
        h1_xml = zf.read('word/header1.xml').decode('utf-8')
        assert 'drawing' in h1_xml or 'shape' in h1_xml
    print("  - Header Logos in header1.xml -> OK")
    
    results['TEST_1_FRONT_MATTER'] = 'PASS' if t1_pass else 'FAIL'
    
    # -------------------------------------------------------------------------
    # TEST 2: Page Margins Compliance
    # -------------------------------------------------------------------------
    print("\n[TEST 2] Page Margins Compliance")
    t2_pass = True
    
    sectPr = body_elems[-1] if body_elems[-1].tag.endswith('sectPr') else body.find('.//w:sectPr', NS)
    pgMar = sectPr.find('w:pgMar', NS)
    assert pgMar is not None, "pgMar not found in sectPr!"
    
    top = int(pgMar.get(f'{{{W_NS}}}top'))
    bottom = int(pgMar.get(f'{{{W_NS}}}bottom'))
    left = int(pgMar.get(f'{{{W_NS}}}left'))
    right = int(pgMar.get(f'{{{W_NS}}}right'))
    
    print(f"  Top Margin:    {top} twips ({top/1440:.2f} in) [Expected: 2301 twips / 1.60 in]")
    print(f"  Left Margin:   {left} twips ({left/1440:.2f} in) [Expected: 2301 twips / 1.60 in]")
    print(f"  Bottom Margin: {bottom} twips ({bottom/1440:.2f} in) [Expected: 1729 twips / 1.20 in]")
    print(f"  Right Margin:  {right} twips ({right/1440:.2f} in) [Expected: 2017-2018 twips / 1.40 in]")
    
    assert top == 2301, f"Top margin mismatch: {top}"
    assert left == 2301, f"Left margin mismatch: {left}"
    assert bottom == 1729, f"Bottom margin mismatch: {bottom}"
    assert right in [2017, 2018], f"Right margin mismatch: {right}"
    print("  - Margins exactly conform to institutional template and Thesis p.15 -> OK")
    results['TEST_2_MARGINS'] = 'PASS'

    # -------------------------------------------------------------------------
    # TEST 3: Typography & Spacing Properties
    # -------------------------------------------------------------------------
    print("\n[TEST 3] Typography & Spacing Properties")
    t3_pass = True
    
    h1_count = 0
    h2_count = 0
    body_count = 0
    
    for idx in range(28, len(doc.paragraphs)):
        p = doc.paragraphs[idx]
        txt = p.text.strip()
        if not txt:
            continue
            
        is_h1 = (
            txt.startswith("CHAPTER") or
            txt in ["INTRODUCTION", "REVIEW OF RELATED LITERATURE AND THEORETICAL/CONCEPTUAL FRAMEWORK", "REFERENCES"] or
            any(txt.startswith(f"1.{n} ") for n in range(1, 7)) or
            any(txt.startswith(f"2.{n} ") for n in range(1, 6))
        )
        is_h2 = (
            any(txt.startswith(f"1.{n}.{m} ") for n in range(1, 7) for m in range(1, 5)) or
            any(txt.startswith(f"2.{n}.{m} ") for n in range(1, 6) for m in range(1, 5))
        )
        
        if is_h1:
            h1_count += 1
            assert p.paragraph_format.keep_with_next == True, f"H1 at P{idx} keep_with_next is False!"
            for r in p.runs:
                if r.text.strip():
                    assert r.font.name == "Arial", f"H1 at P{idx} font is {r.font.name}"
                    assert r.bold == True, f"H1 at P{idx} is not bold"
                    assert r.font.size.pt == 12.0, f"H1 at P{idx} size is {r.font.size.pt}"
            sp_before_pt = p.paragraph_format.space_before.pt if p.paragraph_format.space_before else 0
            if txt not in ["INTRODUCTION", "REVIEW OF RELATED LITERATURE AND THEORETICAL/CONCEPTUAL FRAMEWORK"]:
                assert sp_before_pt >= 12.0, f"H1 at P{idx} space_before too small: {sp_before_pt}pt"
        elif is_h2:
            h2_count += 1
            assert p.paragraph_format.keep_with_next == True, f"H2 at P{idx} keep_with_next is False!"
            for r in p.runs:
                if r.text.strip():
                    assert r.font.name == "Arial", f"H2 at P{idx} font is {r.font.name}"
                    assert r.bold == True, f"H2 at P{idx} is not bold"
                    assert r.font.size.pt == 11.0, f"H2 at P{idx} size is {r.font.size.pt}"
            sp_before_pt = p.paragraph_format.space_before.pt if p.paragraph_format.space_before else 0
            assert sp_before_pt >= 8.0, f"H2 at P{idx} space_before too small: {sp_before_pt}pt"
        elif idx < 155 and not txt.startswith("Figure") and not txt.startswith("Table") and not txt.startswith("Adapted") and not txt.startswith("Input-Process") and not txt.startswith("Quantitative"):
            pPr = p._element.pPr
            ind = pPr.find(f'{{{W_NS}}}ind') if pPr is not None else None
            ind_first = int(ind.get(f'{{{W_NS}}}firstLine', 0)) if ind is not None and ind.get(f'{{{W_NS}}}firstLine') else 0
            
            if ind_first == 720:
                body_count += 1
                assert p.paragraph_format.line_spacing == 1.5, f"Body P{idx} line spacing {p.paragraph_format.line_spacing}"
                if idx not in [130, 143, 149]: # P130 follows Table 1, P143 follows Figure 1, P149 follows Figure 2
                    assert p.paragraph_format.space_before.pt == 0, f"Body P{idx} space_before {p.paragraph_format.space_before.pt}"
                assert p.paragraph_format.space_after.pt == 0, f"Body P{idx} space_after {p.paragraph_format.space_after.pt}"
                assert p.paragraph_format.keep_with_next == False, f"Body P{idx} keep_with_next is True!"
                for r in p.runs:
                    if r.text.strip():
                        assert r.font.name == "Arial", f"Body P{idx} font is {r.font.name}"
                        assert r.font.size.pt == 11.0, f"Body P{idx} size is {r.font.size.pt}"

    print(f"  - Heading 1 instances verified: {h1_count} (Arial 12pt Bold, keep_with_next=True, space_before ~14pt)")
    print(f"  - Heading 2 instances verified: {h2_count} (Arial 11pt Bold, keep_with_next=True, space_before ~10pt)")
    print(f"  - Body narrative paragraphs verified: {body_count} (Arial 11pt Regular, 1.5 line spacing, 0.5\" indent, 0pt space)")
    results['TEST_3_TYPOGRAPHY'] = 'PASS'

    # -------------------------------------------------------------------------
    # TEST 4: Table 1 Synthesis Matrix Structure
    # -------------------------------------------------------------------------
    print("\n[TEST 4] Table 1 Synthesis Matrix Structure")
    assert len(doc.tables) >= 1, "No tables found!"
    t = doc.tables[0]
    assert len(t.rows) == 11, f"Table rows count: {len(t.rows)}"
    assert len(t.columns) == 6, f"Table cols count: {len(t.columns)}"
    print(f"  - Dimensions: {len(t.rows)} rows x {len(t.columns)} columns -> OK")
    
    tblPr = t._element.tblPr
    tblBorders = tblPr.find(f'{{{W_NS}}}tblBorders')
    assert tblBorders is not None, "tblBorders missing!"
    for b_name in ['top', 'left', 'bottom', 'right', 'insideH', 'insideV']:
        b = tblBorders.find(f'{{{W_NS}}}{b_name}')
        assert b is not None, f"Border {b_name} missing!"
        assert b.get(f'{{{W_NS}}}val') == 'single', f"Border {b_name} val not single"
        assert b.get(f'{{{W_NS}}}sz') == '6', f"Border {b_name} sz not 6 (0.75pt)"
        assert b.get(f'{{{W_NS}}}color').upper() == '444444', f"Border {b_name} color not 444444"
    print("  - Borders: single 0.75pt (sz=\"6\") color #444444 on all borders -> OK")
    
    r0_trPr = t.rows[0]._element.trPr
    assert r0_trPr.find(f'{{{W_NS}}}tblHeader') is not None, "Row 0 missing <w:tblHeader/>"
    print("  - Row 0 <w:tblHeader/> (repeating header row) -> OK")
    
    for r_idx, r in enumerate(t.rows):
        trPr = r._element.trPr
        assert trPr.find(f'{{{W_NS}}}cantSplit') is not None, f"Row {r_idx} missing <w:cantSplit/>"
    print("  - <w:cantSplit/> present on all 11 rows -> OK")
    
    for c_idx, c in enumerate(t.rows[0].cells):
        txt = c.text.strip()
        assert txt == txt.upper(), f"Header cell {c_idx} not uppercase: {txt}"
        for cp in c.paragraphs:
            assert cp.alignment == WD_ALIGN_PARAGRAPH.CENTER, f"Header cell {c_idx} not centered"
            for cr in cp.runs:
                if cr.text.strip():
                    assert cr.bold == True, f"Header cell {c_idx} run not bold"
    print("  - Header cells: Uppercase, Bold, Centered Alignment -> OK")
    results['TEST_4_TABLE_MATRIX'] = 'PASS'

    # -------------------------------------------------------------------------
    # TEST 5: Embedded Figures & Captions
    # -------------------------------------------------------------------------
    print("\n[TEST 5] Embedded Figures & Captions")
    assert len(media_files) >= 2, f"Not enough media files: {media_files}"
    print(f"  - Embedded media files inside word/media/: {media_files}")
    
    p140 = doc.paragraphs[140]
    assert p140.text.strip() == "Figure 1", f"P140 text: {p140.text}"
    assert p140.runs[0].bold == True, "Figure 1 label not bold"
    
    p141 = doc.paragraphs[141]
    assert "DeLone and McLean" in p141.text and "2003" in p141.text, f"P141 text: {p141.text}"
    assert p141.runs[0].italic == True, "Figure 1 title not italic"
    
    p142 = doc.paragraphs[142]
    assert '<w:drawing' in p142._element.xml, "Figure 1 image drawing missing"
    assert p142.alignment == WD_ALIGN_PARAGRAPH.CENTER, "Figure 1 image not centered"
    
    p143 = doc.paragraphs[143]
    assert "Figure 1 illustrates the theoretical framework" in p143.text
    print("  - Figure 1: Bold label, Italic title above, Centered inline image, Narrative below -> OK")
    
    p146 = doc.paragraphs[146]
    assert p146.text.strip() == "Figure 2", f"P146 text: {p146.text}"
    assert p146.runs[0].bold == True, "Figure 2 label not bold"
    
    p147 = doc.paragraphs[147]
    assert "Input-Process-System-Output (IPSO)" in p147.text, f"P147 text: {p147.text}"
    assert p147.runs[0].italic == True, "Figure 2 title not italic"
    
    p148 = doc.paragraphs[148]
    assert '<w:drawing' in p148._element.xml, "Figure 2 image drawing missing"
    assert p148.alignment == WD_ALIGN_PARAGRAPH.CENTER, "Figure 2 image not centered"
    
    p149 = doc.paragraphs[149]
    assert "Figure 2 outlines the conceptual architecture" in p149.text
    print("  - Figure 2: Bold label, Italic title above, Centered inline image, Narrative below -> OK")
    results['TEST_5_FIGURES'] = 'PASS'

    # -------------------------------------------------------------------------
    # TEST 6: References Layout
    # -------------------------------------------------------------------------
    print("\n[TEST 6] References Layout")
    p154 = doc.paragraphs[154]
    assert p154.text.strip() == "2.5 References", f"P154 text: {p154.text}"
    
    ref_entries = doc.paragraphs[156:170]
    assert len(ref_entries) == 14, f"Reference entries count: {len(ref_entries)}"
    print(f"  - Total reference entries: {len(ref_entries)}")
    
    for r_i, rp in enumerate(ref_entries):
        txt = rp.text.strip()
        pPr = rp._element.pPr
        ind = pPr.find(f'{{{W_NS}}}ind')
        left = int(ind.get(f'{{{W_NS}}}left', 0))
        hanging = int(ind.get(f'{{{W_NS}}}hanging', 0))
        first_line = int(ind.get(f'{{{W_NS}}}firstLine', 0))
        
        assert left == 720 and (hanging == 720 or first_line == -720), f"Ref {r_i} indent: left={left}, hanging={hanging}, first={first_line}"
        assert rp.paragraph_format.line_spacing == 1.5, f"Ref {r_i} line spacing: {rp.paragraph_format.line_spacing}"
        for r in rp.runs:
            if r.text.strip():
                assert r.font.name == "Arial", f"Ref {r_i} font is {r.font.name}"
                assert r.font.size.pt == 11.0, f"Ref {r_i} size is {r.font.size.pt}"
                
    print("  - All 14 reference entries verified (0.5\" hanging indent [left=720, hanging=720], 1.5 line spacing, Arial 11pt Regular) -> OK")
    results['TEST_6_REFERENCES'] = 'PASS'

    print("\n" + "=" * 80)
    print("ALL EMPIRICAL TESTS PASSED PERFECTLY!")
    print(f"TEST RESULTS: {results}")
    print("OVERALL VERDICT: APPROVE")
    print("=" * 80)
    return True

if __name__ == '__main__':
    run_suite()
