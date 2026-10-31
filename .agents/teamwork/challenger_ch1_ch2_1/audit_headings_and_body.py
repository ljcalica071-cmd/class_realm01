from docx import Document

TARGET_DOCX = r"c:\Users\Admin\AndroidStudioProjects\class_realm01\INTEGRATED SENIOR HIGH SCHOOL ATTENDANCE.docx"
doc = Document(TARGET_DOCX)

print(f"Total paragraphs: {len(doc.paragraphs)}")

heading_candidates = []
body_narrative = []
list_items = []
figures_and_captions = []
references = []

for i in range(28, len(doc.paragraphs)):
    p = doc.paragraphs[i]
    txt = p.text.strip()
    if not txt:
        continue
    
    pPr = p._element.pPr
    keep_next = pPr.keepNext is not None if pPr is not None else False
    
    sp = pPr.find('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}spacing') if pPr is not None else None
    sp_before = int(sp.get('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}before', 0)) if sp is not None else 0
    sp_after = int(sp.get('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}after', 0)) if sp is not None else 0
    line_sp = int(sp.get('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}line', 0)) if sp is not None else 0
    
    ind = pPr.find('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}ind') if pPr is not None else None
    ind_left = int(ind.get('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}left', 0)) if ind is not None and ind.get('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}left') else 0
    ind_first = int(ind.get('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}firstLine', 0)) if ind is not None and ind.get('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}firstLine') else 0
    ind_hanging = int(ind.get('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}hanging', 0)) if ind is not None and ind.get('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}hanging') else 0

    runs = p.runs
    fonts = set(r.font.name for r in runs if r.text.strip())
    sizes = set(r.font.size.pt for r in runs if r.font.size and r.text.strip())
    bolds = set(r.bold for r in runs if r.text.strip())
    italics = set(r.italic for r in runs if r.text.strip())
    
    info = {
        'idx': i,
        'text': txt,
        'keep_next': keep_next,
        'sp_before': sp_before,
        'sp_after': sp_after,
        'line_sp': line_sp,
        'ind_left': ind_left,
        'ind_first': ind_first,
        'ind_hanging': ind_hanging,
        'fonts': fonts,
        'sizes': sizes,
        'bolds': bolds,
        'italics': italics,
        'alignment': p.alignment
    }
    
    if i >= 156:
        references.append(info)
    elif txt.startswith("Figure 1") or txt.startswith("Figure 2") or txt.startswith("Adapted DeLone") or txt.startswith("Input-Process") or txt.startswith("Table 1") or txt.startswith("Quantitative Synthesis Matrix"):
        figures_and_captions.append(info)
    elif keep_next or txt.startswith("CHAPTER") or any(txt.startswith(h) for h in [
        "1.1", "1.2", "1.3", "1.4", "1.5", "1.6", "2.1", "2.2", "2.3", "2.4", "2.5", "INTRODUCTION", "REVIEW OF"
    ]):
        heading_candidates.append(info)
    elif ind_hanging > 0 or (txt[0].isdigit() and (txt[1] == '.' or txt[2] == '.')) or txt.startswith("•"):
        list_items.append(info)
    else:
        body_narrative.append(info)

print(f"Categories identified:")
print(f"  Headings: {len(heading_candidates)}")
print(f"  Body Narrative Paragraphs: {len(body_narrative)}")
print(f"  List / Definition items: {len(list_items)}")
print(f"  Figures / Captions / Tables: {len(figures_and_captions)}")
print(f"  References: {len(references)}")

print("\n--- Detailed Heading Audit ---")
for h in heading_candidates:
    print(f"P{h['idx']:03d} | '{h['text'][:45]}' | keep_next={h['keep_next']} | sp_before={h['sp_before']} twips ({h['sp_before']/20}pt) | sp_after={h['sp_after']} twips | font={h['fonts']} | size={h['sizes']} | bold={h['bolds']}")

print("\n--- Body Narrative Summary Audit ---")
non_compliant_body = []
for b in body_narrative:
    # check: Arial, 11pt, regular, 1.5 spacing (360), indent firstLine=720, sp_before=0, sp_after=0
    issues = []
    if not (b['fonts'] == {'Arial'} or b['fonts'] == set()):
        issues.append(f"font={b['fonts']}")
    if not (b['sizes'] == {11.0} or b['sizes'] == set()):
        issues.append(f"size={b['sizes']}")
    if True in b['bolds']:
        # check if it has bold runs
        pass
    if b['line_sp'] != 360:
        issues.append(f"line_sp={b['line_sp']}")
    if b['ind_first'] != 720:
        issues.append(f"ind_first={b['ind_first']}")
    if b['sp_before'] != 0:
        issues.append(f"sp_before={b['sp_before']}")
    if b['sp_after'] != 0:
        issues.append(f"sp_after={b['sp_after']}")
    if issues:
        non_compliant_body.append((b['idx'], b['text'][:40], issues))

print(f"Compliant body narrative count: {len(body_narrative) - len(non_compliant_body)} / {len(body_narrative)}")
if non_compliant_body:
    print(f"Non-compliant body paragraphs ({len(non_compliant_body)}):")
    for idx, txt, iss in non_compliant_body[:10]:
        print(f"  P{idx:03d} ('{txt}'): {', '.join(iss)}")
