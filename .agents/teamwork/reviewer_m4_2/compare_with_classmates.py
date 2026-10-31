import os
from docx import Document

def inspect(title, path):
    print("=" * 60)
    print(title, ":", path)
    if not os.path.exists(path):
        print("Not found")
        return
    doc = Document(path)
    print("Paragraph count:", len(doc.paragraphs))
    print("Table count:", len(doc.tables))
    print("Sections count:", len(doc.sections))
    
    print("\nHeadings:")
    for i, p in enumerate(doc.paragraphs):
        t = p.text.strip()
        if t and (t.startswith("PART") or t.startswith("Part") or t.startswith("ACTIVITY") or t.startswith("Activity") or t.startswith("Table") or t.startswith("Figure")):
            print(f"  P{i}: {t[:90]}")
            
    print("\nTables:")
    for i, t in enumerate(doc.tables):
        r_cnt = len(t.rows)
        c_cnt = len(t.columns)
        h = [c.text.strip().replace('\n', ' ') for c in t.rows[0].cells] if r_cnt > 0 else []
        print(f"  T{i+1}: rows={r_cnt}, cols={c_cnt}, headers={h[:5]}")

if __name__ == '__main__':
    inspect("CLASSMATE ACT 4", r"c:\Users\Admin\Downloads\Activity#4_UML_Object.docx")
